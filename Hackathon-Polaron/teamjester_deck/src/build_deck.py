"""Build the seven-slide TeamJester overview deck.

    python src/build_deck.py --repo /path/to/teamjester_main_<sha>
    python src/build_deck.py --repo ... --new-batch new_batch_results/<file>
                             [--new-batch-name "Batch 4"]

Inputs : provenance/deck_numbers.json (src/deck_numbers.py),
         assets/*.png + assets/annotations.json (src/make_assets.py),
         provenance/verify_summary.json (src/verify_main.py).
Outputs: TeamJester_Algorithm_Overview.pptx / .pdf, previews/slide-N.png,
         CLAIM_LEDGER.csv, PRESENTER_SCRIPT.md.
"""
import argparse
import csv
import json
import os
import re
import shutil
import subprocess
import sys

from lxml import etree
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_MARKER_STYLE, XL_TICK_MARK,
                             XL_TICK_LABEL_POSITION)
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import style as S  # noqa: E402
import new_batch as NB  # noqa: E402

OUT_NAME = "TeamJester_Algorithm_Overview"
W, H = 13.333, 7.5
ML = 0.55
CW = W - 2 * ML
N_SLIDES = 7
A = os.path.join(DECK, "assets")


# =============================================================================
# numbers + claim ledger
# =============================================================================
class Nums:
    def __init__(self):
        with open(os.path.join(DECK, "provenance", "deck_numbers.json")) as f:
            self.d = json.load(f)

    def t(self, k):
        return self.d[k]["text"]

    def v(self, k):
        return self.d[k]["value"]

    def pct(self, k, nd=0, signed=False):
        v = self.v(k) * 100
        s = f"{v:+.{nd}f}%" if signed else f"{v:.{nd}f}%"
        return s.replace("-", "−")


def verification_for(source, verify):
    s = source
    if "presentation-derived" in s or s.startswith("derived"):
        return "derived for this deck from main's committed per-particle table"
    if "loo_summary_regenerated" in s:
        return ("committed loo_summary.csv is stale; value from re-running "
                "main's run_all.stage_stats() (src/verify_main.py)")
    if "feature_meta" in s:
        return ("committed count; a from-scratch build of main's current "
                "features.py registers more features (see build-state caveat "
                "in the slide 2/7 notes)")
    if "point_count_validation" in s:
        return ("read from the committed sheet: 300 rows, human_class column "
                "empty (no expert labels yet)")
    if "features.csv + quality_guards.csv" in s:
        return ("feature columns re-computed identical (31/31 masks); "
                "quality_guards values identical in the from-scratch rebuild")
    if any(x in s for x in ("features.csv", "particles.csv", "masks",
                            "exp_newfeat.csv")):
        seg = verify.get("segmentation_pore_identical")
        return (f"re-computed with main's code: masks identical for {seg}/31 "
                f"images, {verify.get('features_matching')}/"
                f"{verify.get('features_compared')} feature columns identical")
    if "deltas" in s:
        return ("regenerated with main's code: deltas, % changes, "
                "Mann-Whitney p and classes identical; bootstrap CI / Holm p "
                "differ slightly (committed table predates si_d10_um removal)")
    if any(x in s for x in ("baseline_stats", "loo_predictions",
                            "images_needed", "repeatability")):
        return "regenerated with main's run_all.stage_stats(): identical"
    if any(x in s for x in ("quality_guards", "segment_info",
                            "threshold_sensitivity", "cross_detector",
                            "half_agreement")):
        st = verify.get("clean_rebuild", {}).get(os.path.basename(
            s.split(" ")[0]))
        return ("from-scratch rebuild with main's code: " + st) if st else \
            "committed output of main's segment.py (not independently re-run)"
    if "dfn" in s:
        return ("committed PyBaMM output; censoring confirmed from dfn.py "
                "(t_end = 1.05 h / C-rate) and the table itself")
    if "TIFF" in s or "Hackathon-Polaron" in s:
        return "read directly from main's dataset files"
    if "git" in s:
        return "git metadata of the pinned worktree"
    if ".py" in s:
        return "read from main's source code at the pinned commit"
    return "committed table at the pinned commit"


class Ledger:
    def __init__(self, nums, verify, sha):
        self.N, self.verify, self.sha = nums, verify, sha
        self.rows = []

    def add(self, slide, claim, *keys, note=""):
        for k in keys or (None,):
            if k is None:
                src, how, val = "see note", note, ""
            else:
                e = self.N.d[k]
                src, how, val = e["source"], e["how"], e["text"] or \
                    json.dumps(e["value"])[:120]
            self.rows.append(dict(
                id=f"S{slide}-{len(self.rows) + 1:03d}", slide=slide,
                claim=claim, number_key=k or "", value=val,
                source=f"{src} @ {self.sha[:7]}" if k else src,
                how_obtained=how, verification=verification_for(
                    src, self.verify) if k else note))
        return claim

    def write(self, path):
        with open(path, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(self.rows[0]))
            w.writeheader()
            w.writerows(self.rows)
        md = [f"# Claim-to-source ledger — TeamJester overview deck", "",
              f"Every number on the slides, its source file in "
              f"[TeamJester main @ {self.sha[:7]}]"
              f"({self.N.v('repo_url')}/tree/{self.sha}) and how it was "
              "verified. Machine-readable version: `" +
              os.path.basename(path) + "`.", ""]
        for sl in sorted({r["slide"] for r in self.rows}):
            md += [f"## Slide {sl}", "",
                   "| id | claim | value | source | how obtained | "
                   "verification |", "|---|---|---|---|---|---|"]
            for r in self.rows:
                if r["slide"] == sl:
                    cells = [str(r[k]).replace("|", "/").replace("\n", " ")
                             for k in ("id", "claim", "value", "source",
                                       "how_obtained", "verification")]
                    md.append("| " + " | ".join(cells) + " |")
            md.append("")
        with open(path[:-4] + ".md", "w") as f:
            f.write("\n".join(md))


# =============================================================================
# drawing helpers
# =============================================================================
def rgb(h):
    return RGBColor.from_string(h)


def strip_style(shape):
    """Drop the theme style reference (LibreOffice draws its effect as a
    drop shadow); fill/line/text colours are always set explicitly here."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def tw(s, pt, bold=False):
    """Approximate rendered width (in) of a string in LM Sans 10."""
    return len(s) * pt * (0.49 if bold else 0.45) / 72.0


def n_lines(s, pt, width_in, bold=False):
    import math
    return sum(max(1, math.ceil(tw(part, pt, bold) / max(width_in, 0.1)))
               for part in s.split("\n"))


def run_fmt(r, font=S.BODY_FONT, size=20, bold=False, italic=False,
            color=S.NAVY, link=None, sup=False):
    f = r.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)
    if sup:
        rPr.set("baseline", "30000")
    if link:
        r.hyperlink.address = link
        # keep our colour instead of the theme hyperlink blue
        r.font.color.rgb = rgb(color)


def text(slide, x, y, w, h, paras, anchor="t", margin=(0.0, 0.0, 0.0, 0.0),
         wrap=True, name=None):
    """paras: list of dicts {runs:[(text, fmt)], align, sb, sa, ls, bullet}"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (
        Inches(m) for m in margin)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, p in enumerate(paras):
        if isinstance(p, tuple):
            p = dict(runs=[p])
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                          "r": PP_ALIGN.RIGHT}[p.get("align", "l")]
        if p.get("sb") is not None:
            para.space_before = Pt(p["sb"])
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        if p.get("bullet"):
            add_bullet(para, p)
        for t_, fmt in p["runs"]:
            r = para.add_run()
            r.text = t_
            run_fmt(r, **fmt)
    return box


def F(size=20, **k):
    return dict(size=size, **k)


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, dash=None,
         shape=MSO_SHAPE.RECTANGLE, radius=None, name=None, alpha=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w),
                               Inches(h))
    if name:
        s.name = name
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            sf = s.fill._xPr.find(qn("a:solidFill"))
            clr = sf.find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val",
                                                     str(int(alpha * 100000)))
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = dash
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    strip_style(s)
    s.text_frame.text = ""
    return s


def line(slide, x1, y1, x2, y2, color=S.NAVY, lw=1.25, dash=None,
         head=None, tail=None, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                   Inches(y1), Inches(x2), Inches(y2))
    if name:
        c.name = name
    strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = dash
    ln = c.line._get_or_add_ln()
    for tag, kind in (("a:headEnd", head), ("a:tailEnd", tail)):
        if kind:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", kind)
            e.set("w", "med")
            e.set("len", "med")
    return c


def arrow(slide, x1, y1, x2, y2, color=S.NAVY, lw=1.5, name=None):
    return line(slide, x1, y1, x2, y2, color, lw, tail="triangle", name=name)


def pic(slide, path, x, y, w=None, h=None, name=None):
    p = slide.shapes.add_picture(path, Inches(x), Inches(y),
                                 Inches(w) if w else None,
                                 Inches(h) if h else None)
    if name:
        p.name = name
    return p


def freeform(slide, pts_in, color=S.TEAL, lw=2.0, close=True, name=None):
    ff = slide.shapes.build_freeform(Inches(pts_in[0][0]),
                                     Inches(pts_in[0][1]), scale=1.0)
    ff.add_line_segments([(Inches(x), Inches(y)) for x, y in pts_in[1:]],
                         close=close)
    s = ff.convert_to_shape()
    if name:
        s.name = name
    strip_style(s)
    s.fill.background()
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def label(slide, x, y, w, h, t_, size=13, color=S.NAVY, fill=S.WHITE,
          bold=False, align="l", alpha=0.92, line_color=None, name=None):
    r = rect(slide, x, y, w, h, fill=fill, line=line_color, lw=0.5,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18, alpha=alpha,
             name=name)
    tf = r.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                   "r": PP_ALIGN.RIGHT}[align]
    rr = p.add_run()
    rr.text = t_
    run_fmt(rr, size=size, color=color, bold=bold)
    return r


def scalebar(slide, x, y, length_in, txt, color=S.WHITE, bg=S.NAVY,
             size=12):
    rect(slide, x - 0.08, y - 0.30, max(length_in, 0.7) + 0.16, 0.42, fill=bg,
         alpha=0.72, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2,
         name="scale bar box")
    rect(slide, x, y, length_in, 0.055, fill=color, name="scale bar")
    text(slide, x, y - 0.31, max(length_in, 0.7), 0.28,
         [dict(runs=[(txt, F(size, color=color))], align="c")],
         name="scale bar label")


def title(slide, t_, sub=None):
    text(slide, ML, 0.30, CW, 0.72,
         [dict(runs=[(t_, F(34, font=S.TITLE_FONT))])], anchor="b",
         name="Title")
    line(slide, ML, 1.06, ML + 1.1, 1.06, color=S.TEAL, lw=2.25,
         name="title rule")
    if sub:
        text(slide, ML, 1.12, CW, 0.5,
             [dict(runs=[(sub, F(20, color=S.NAVY2))])], name="Lead")


def footer(slide, n, N):
    text(slide, ML, 7.08, 9.5, 0.3,
         [dict(runs=[(f"TeamJester · GitHub main @ {N.t('sha7')} "
                      f"({N.t('commit_date')}) · explainable SEM batch "
                      "comparison", F(10.5, color=S.GREY))])],
         name="footer")
    text(slide, W - ML - 1.5, 7.08, 1.5, 0.3,
         [dict(runs=[(f"{n} / {N_SLIDES}", F(10.5, color=S.GREY))],
               align="r")], name="slide number")


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


def set_table_style_none(tbl):
    tblPr = tbl._tbl.tblPr
    for k in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol",
              "bandCol"):
        tblPr.set(k, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid


def cell_borders(cell, top=None, bottom=None):
    """booktabs rules: (width_pt, colour) or None = no line."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    spec = {"a:lnL": None, "a:lnR": None, "a:lnT": top, "a:lnB": bottom}
    for i, (tag, sp) in enumerate(spec.items()):
        ln = etree.Element(qn(tag))
        if sp is None:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        else:
            ln.set("w", str(int(Pt(sp[0]))))
            ln.set("cap", "flat")
            ln.set("cmpd", "sng")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", sp[1])
            etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        tcPr.insert(i, ln)


def cell_text(cell, paras, fill=None, margins=(0.07, 0.05, 0.07, 0.05),
              anchor="t"):
    cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom \
        = (Inches(m) for m in margins)
    cell.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[anchor]
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
    else:
        cell.fill.background()
    tf = cell.text_frame
    tf.word_wrap = True
    for i, p in enumerate(paras):
        if isinstance(p, tuple):
            p = dict(runs=[p])
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if p.get("sb") is not None:
            para.space_before = Pt(p["sb"])
        if p.get("ls") is not None:
            para.line_spacing = p["ls"]
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER}[
            p.get("align", "l")]
        if p.get("bullet"):
            add_bullet(para, p)
        for t_, fmt in p["runs"]:
            r = para.add_run()
            r.text = t_
            run_fmt(r, **fmt)


def add_bullet(para, p):
    pPr = para._p.get_or_add_pPr()
    ind = p.get("indent", 0.2)
    pPr.set("marL", str(int(Inches(ind))))
    pPr.set("indent", str(-int(Inches(ind))))
    bc = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bc, qn("a:srgbClr")).set("val", p.get("bullet_color",
                                                           S.TEAL))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", S.BODY_FONT)
    etree.SubElement(pPr, qn("a:buChar")).set("char", p["bullet"])


def booktabs(slide, x, y, colw, rowh, header, body, size=13, head_size=13.5,
             name="table"):
    nr, nc = len(body) + 1, len(colw)
    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y),
                                Inches(sum(colw)), Inches(sum(rowh)))
    gf.name = name
    tbl = gf.table
    set_table_style_none(tbl)
    for j, wj in enumerate(colw):
        tbl.columns[j].width = Inches(wj)
    for i, hi in enumerate(rowh):
        tbl.rows[i].height = Inches(hi)
    for j, htxt in enumerate(header):
        c = tbl.cell(0, j)
        cell_text(c, [dict(runs=[(htxt, F(head_size, bold=True,
                                          color=S.NAVY))])], anchor="m")
        cell_borders(c, top=(1.5, S.NAVY), bottom=(0.75, S.NAVY))
    for i, row in enumerate(body, start=1):
        last = i == nr - 1
        for j, paras in enumerate(row):
            c = tbl.cell(i, j)
            cell_text(c, paras)
            cell_borders(c, top=None, bottom=(1.5, S.NAVY) if last else
                         (0.5, S.RULE))
    return tbl


def dot(slide, x, y, d, color, name=None):
    return rect(slide, x, y, d, d, fill=color, shape=MSO_SHAPE.OVAL, name=name)


def repo_link(N, path, pattern=None):
    url = f"{N.v('repo_url')}/blob/{N.v('sha')}/{path}"
    if pattern and REPO:
        with open(os.path.join(REPO, path)) as f:
            for i, ln in enumerate(f, 1):
                if re.search(pattern, ln):
                    return f"{url}#L{i}"
    return url


DOI = {
    "singh": ("Singh, Kaiser & Hahn (2016) Effect of porosity on the thick "
              "electrodes for high energy density lithium ion batteries for "
              "stationary applications. Batteries 2, 35.",
              "https://doi.org/10.3390/batteries2040035"),
    "obrovac": ("Obrovac & Chevrier (2014) Alloy negative electrodes for "
                "Li-ion batteries. Chem. Rev. 114, 11444-11502.",
                "https://doi.org/10.1021/cr500207g"),
    "liu12": ("Liu, X. H. et al. (2012) Size-dependent fracture of silicon "
              "nanoparticles during lithiation. ACS Nano 6, 1522-1531.",
              "https://doi.org/10.1021/nn204476h"),
    "wucui": ("Wu & Cui (2012) Designing nanostructured Si anodes for high "
              "energy lithium ion batteries. Nano Today 7, 414-429.",
              "https://doi.org/10.1016/j.nantod.2012.08.004"),
    "yolk": ("Liu, N. et al. (2012) A yolk-shell design for stabilized and "
             "scalable Li-ion battery alloy anodes. Nano Lett. 12, 3315-3321.",
             "https://doi.org/10.1021/nl3014814"),
    "ebner": ("Ebner, Chung, Garcia & Wood (2014) Tortuosity anisotropy in "
              "lithium-ion battery electrodes. Adv. Energy Mater. 4, 1301278.",
              "https://doi.org/10.1002/aenm.201301278"),
    "billaud": ("Billaud et al. (2016) Magnetically aligned graphite "
                "electrodes for high-rate performance Li-ion batteries. "
                "Nature Energy 1, 16097.",
                "https://doi.org/10.1038/nenergy.2016.97"),
    "otsu": ("Otsu (1979) A threshold selection method from gray-level "
             "histograms. IEEE Trans. SMC 9, 62-66.",
             "https://doi.org/10.1109/TSMC.1979.4310076"),
    "clark": ("Clark & Evans (1954) Distance to nearest neighbor as a "
              "measure of spatial relationships in populations. Ecology 35, "
              "445-453.", "https://doi.org/10.2307/1931034"),
    "chen": ("Chen et al. (2020) Development of experimental techniques for "
             "parameterization of multi-scale lithium-ion battery models. "
             "J. Electrochem. Soc. 167, 080534.",
             "https://doi.org/10.1149/1945-7111/ab9050"),
    "pybamm": ("Sulzer et al. (2021) Python Battery Mathematical Modelling "
               "(PyBaMM). J. Open Res. Softw. 9, 14.",
               "https://doi.org/10.5334/jors.309"),
    "holm": ("Holm (1979) A simple sequentially rejective multiple test "
             "procedure. Scand. J. Statist. 6, 65-70.",
             "https://www.jstor.org/stable/4615733"),
}
REPO = None
CR = {}


def rebuild_caveat(N):
    if not CR:
        return ""
    extra = len(CR["kept_fresh"]) - len(CR["kept_committed"])
    hs = "; ".join(f"{b_.replace('_', ' ')} {f_}" for b_, f_ in
                   CR["holm_survivors_fresh"])
    return (
        "\n\nBUILD-STATE CAVEAT (from-scratch rebuild of main's code, "
        "provenance/clean_rebuild_summary.json)\n"
        "Re-running run_all.py at this commit on an empty outputs/ folder "
        "reproduces every segmentation table and all committed feature "
        "columns exactly, and every % shift and class of the "
        f"{len(CR['kept_committed'])} committed markers. However, the current "
        "features.py also computes "
        f"{len(CR['extra_columns'])} newer columns that the committed "
        f"features.csv lacks; {extra} of them pass main's gates, so a fresh "
        f"build keeps {len(CR['kept_fresh'])} markers instead of "
        f"{len(CR['kept_committed'])}. The classifier numbers then change: "
        f"all-features {CR['loo_allfeatures_fresh'][0]}/"
        f"{CR['loo_allfeatures_fresh'][1]} correct (committed "
        f"{N.t('loo_correct')}/31), artefact-safe "
        f"{CR['loo_safe_fresh'][0]}/{CR['loo_safe_fresh'][1]} (committed "
        f"13/31), held-out Batch 3 called OUT {CR['b3_out_fresh'][0]}/"
        f"{CR['b3_out_fresh'][1]}, shuffled-label control "
        f"{100 * CR['shuffled_fresh']:.0f}%; Holm survivors become: {hs}. "
        "The deck reports the committed run (the tables main's webapp loads); treat "
        "classifier accuracy as UNRESOLVED — it depends on the build state — "
        "while the structural markers are robust to it.")


def provenance_block(N):
    return (
        "\n\nPROVENANCE (applies to every number on this slide)\n"
        f"Repository: {N.v('repo_url')} (private; team access)\n"
        f"Branch/commit: main @ {N.v('sha')} ({N.t('commit_date')}), "
        "inspected in a clean detached worktree; nothing from other branches, "
        "archived methods or local experiments is used.\n"
        "Numbers: computed by src/deck_numbers.py from main's committed tables "
        "(analysis/outputs/tables/*); each one is listed with file, column and "
        "verification status in CLAIM_LEDGER.csv.\n"
        "Verification: src/verify_main.py re-ran main's own code in a Python "
        "3.10 environment matching main's .devin/blueprint.yaml — all 31 "
        "segmentations and all 54 committed feature columns reproduce exactly; "
        "baseline and leave-one-out tables reproduce exactly.")


# =============================================================================
# SLIDES
# =============================================================================
def slide1(prs, N, L, ann):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Do new anode batches match the approved one?",
          "An explainable image-analysis pipeline for battery-electrode "
          "microscope photos")
    hx, hy, hw = ML, 1.66, CW
    a = ann["hero"]
    hh = hw * a["h"] / a["w"]
    pic(s, os.path.join(A, "hero.png"), hx, hy, hw, hh, name="hero image")
    fx = lambda px: hx + hw * px / a["w"]          # noqa: E731
    fy = lambda py: hy + hh * py / a["h"]          # noqa: E731
    # callouts: label in the top or bottom band of the photo, arrow to target
    callouts = [("bright particle — assumed silicon-based", a["particle"],
                 "top", -0.45),
                ("grey flake — assumed graphite", a["graphite"], "top", -0.2),
                ("black gap — pore or crack", a["pore"], "bottom", 0.15)]
    for t_, (px, py), band, shift in callouts:
        tx, ty = fx(px), fy(py)
        lw_ = tw(t_, 13.5) + 0.3
        lx = min(max(tx - lw_ / 2 + shift * lw_, hx + 0.12), hx + hw - lw_ -
                 0.12)
        ly = hy + 0.14 if band == "top" else hy + hh - 0.5
        ax_ = min(max(tx, lx + 0.3), lx + lw_ - 0.3)
        arrow(s, ax_, ly + (0.34 if band == "top" else 0), tx, ty,
              color=S.WHITE, lw=2.0, name="callout arrow")
        label(s, lx, ly, lw_, 0.34, t_, size=13.5, name="callout")
    sb_um = 20
    sbl = hw * sb_um / 0.025 / a["w"]
    scalebar(s, hx + hw - sbl - 0.3, hy + hh - 0.16, sbl, f"{sb_um} μm")
    text(s, hx, hy + hh + 0.03, hw, 0.28, [dict(runs=[(
        f"{N.t('ex_batch')} · img_{N.t('ex_id')} · BSE detector · "
        f"{N.t('nm_per_px')} per pixel · brightness stretched for display "
        "(same 1st–99.5th percentile stretch the pipeline uses)",
        F(11, color=S.GREY))], align="r")], name="image id")
    L.add(1, "Example photo is Batch 3 image cfe5vt7s (BSE), 25 nm per pixel",
          "ex_id", "nm_per_px")

    by, gap = hy + hh + 0.42, 0.28
    bw = (CW - 3 * gap) / 4
    blocks = [
        ("THE PROBLEM", "Does a new delivery of anode coating match the "
         "approved Batch 3?"),
        ("THE DATA", f"{N.t('n_images')} cross-section photos "
         f"({N.t('n_Batch_1')} + {N.t('n_Batch_2')} + "
         f"{N.t('n_Batch_3')} locations), 3 detectors each"),
        ("WHAT IT PRODUCES", f"Pixel maps, {N.t('n_kept')} repeatable "
         "measurements, shifts vs Batch 3 with uncertainty"),
        ("WHO COULD USE IT", "Quality and process engineers, and suppliers "
         "deciding what to re-check"),
    ]
    L.add(1, "31 photos: 7 + 7 + 17 locations, each with 3 detectors",
          "n_images", "n_Batch_1", "n_Batch_2", "n_Batch_3", "n_tif")
    L.add(1, "20 repeatable measurements retained per photo", "n_kept")
    for i, (h_, b_) in enumerate(blocks):
        x = ML + i * (bw + gap)
        line(s, x, by, x + bw, by, color=S.TEAL, lw=1.0)
        text(s, x, by + 0.05, bw, 0.3,
             [dict(runs=[(h_, F(12.5, bold=True, color=S.TEAL))])])
        text(s, x, by + 0.34, bw, 1.1,
             [dict(runs=[(b_, F(18.5))], ls=0.93)])
    cy = 6.1
    rect(s, ML, cy, CW, 0.86, fill=S.PANEL, name="capability strip")
    text(s, ML + 0.15, cy, 1.5, 0.86,
         [dict(runs=[("What it can\ndo today", F(13.5, bold=True,
                                               color=S.NAVY2))], ls=0.9)],
         anchor="m")
    caps = [
        (S.TEAL, "Image comparison: works",
         "measured shifts come with intervals"),
        (S.AMBER, "Anomaly / batch call: not yet",
         f"single photos: {N.t('loo_correct')} of {N.t('n_images')} "
         "correct"),
        (S.GREY, "Battery simulation: illustrative",
         "optional, model-dependent branch"),
    ]
    L.add(1, "Single-photo batch call: 15 of 31 held-out photos correct",
          "loo_correct", "n_images")
    cx0 = ML + 1.75
    cwid = (ML + CW - cx0) / 3
    for i, (col, h_, b_) in enumerate(caps):
        x = cx0 + i * cwid
        dot(s, x, cy + 0.2, 0.17, col)
        text(s, x + 0.26, cy + 0.08, cwid - 0.32, 0.74,
             [dict(runs=[(h_, F(14, bold=True))]),
              dict(runs=[(b_, F(13, color=S.NAVY2))], sb=1)])
    footer(s, 1, N)
    notes(s, f"""
WHAT TO SAY
This project asks a practical quality question: when a supplier delivers a new batch of anode coating for lithium-ion cells, does it look like the batch we already approved (Batch 3)? Today that judgement is made by eye from microscope photos. The TeamJester pipeline on GitHub main turns those photos into measured, explainable numbers and compares them with Batch 3. The photo is a real cross-section from Batch 3: dark-grey flakes, bright particles and black gaps. Be clear about what works today: comparing structure with uncertainty works; calling a batch from a single photo does not; the battery simulation is an optional illustration.

TERMS
- Anode: the negative electrode; here a coating of graphite flakes with silicon-containing particles (assumed from brightness).
- SEM / BSE: scanning electron microscope; the back-scattered-electron (BSE) detector shows heavier material brighter.
- Cross-section: the coating was cut and polished so we look at its inside.
- Batch 3 = "baseline": the supplier's promised reference; Batch 1 and Batch 2 are later deliveries.

HOW THE NUMBERS WERE OBTAINED
- {N.t('n_images')} photos = rows with subset == full in analysis/outputs/tables/features.csv ({N.t('n_Batch_1')} Batch 1, {N.t('n_Batch_2')} Batch 2, {N.t('n_Batch_3')} Batch 3); {N.t('n_tif')} TIFF files = BSE + Inlens + ETD/SE per location ({N.t('n_etd')} ETD, {N.t('n_se')} SE).
- {N.t('nm_per_px')} per pixel: TIFF XResolution tags (segment_info.csv, um_per_px). Frames are {N.t('fov_w_um')} wide and {N.t('fov_h_um')} tall.
- {N.t('n_kept')} repeatable measurements: rows of baseline_stats.csv (out of {N.t('n_features')} computed; the pruning rule is on slide 2).
- "{N.t('loo_correct')} of {N.t('n_images')} correct": leave-one-out predictions in loo_predictions_allfeatures.csv (every photo held out once). Always answering "Batch 3" would score {N.t('majority_acc')} (17/31) — so single-photo calls are not useful yet.

ASSUMPTIONS AND LIMITS
- "Bright = silicon-based, grey = graphite" is inferred from brightness only (main states this in every output); no chemical analysis (EDS) was done.
- The image is contrast-stretched for display with the same 1st-99.5th percentile stretch the pipeline uses; nothing else is altered.
- The photos are curated locations; we do not know how they were sampled from each batch.

SOURCES
- Repo README and analysis/NOTES.md at main @ {N.v('sha')}: {repo_link(N, 'README.md')}
- Photo: Hackathon-Polaron/Batch_3/img_{N.t('ex_id')}_BSE.tif (rows {ann['hero']['rows'][0]}-{ann['hero']['rows'][1]} shown).

LIKELY QUESTION
Q: "So can it tell me whether a batch is good?"
A: Not on its own. It can tell you, with stated uncertainty, how a batch's measured structure differs from Batch 3, and it flags photos too poor to trust. It does not measure chemistry or performance, and its batch classifier is not yet reliable on single photos.
{provenance_block(N)}""")
    return s


def slide2(prs, N, L):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "How the algorithm works",
          "One fixed recipe, applied identically to every photo — six steps "
          "from picture to comparison")
    cards = [
        ("Read the images", "Three detectors photograph the same spot",
         "thumb_detectors.png", "BSE · Inlens · ETD"),
        ("Label the pixels", "Every pixel becomes pore, grey bulk or bright "
         "particle", "thumb_segment.png", "cut-offs from each photo"),
        ("Check quality", "Is the photo contrasty and clean enough to trust?",
         "thumb_quality.png", f"{N.t('n_flagged')} of {N.t('n_images')} "
         "photos flagged"),
        ("Measure", "Turn the map into numbers; keep only repeatable ones",
         "thumb_measure.png", f"{N.t('n_features')} computed → "
         f"{N.t('n_kept')} kept"),
        ("Compare with Batch 3", "Is a shift bigger than normal photo-to-"
         "photo scatter?", "thumb_compare.png", "interval + corrected test"),
        ("Classify a sample", "Which batch does a new sample resemble most?",
         "thumb_categorise.png", f"single photos: {N.t('loo_correct')}/"
         f"{N.t('n_images')} right"),
    ]
    L.add(2, "2 of 31 photos flagged by the quality guards", "n_flagged",
          "flagged_ids")
    L.add(2, "53 features computed, 20 kept", "n_features", "n_kept")
    L.add(2, "Single photos: 15 of 31 assigned to the right batch",
          "loo_correct")
    gap, cw_ = 0.14, (CW - 5 * 0.14) / 6
    y0, ch = 1.78, 3.32
    for i, (t_, e_, img, cap) in enumerate(cards):
        x = ML + i * (cw_ + gap)
        rect(s, x, y0, cw_, ch, fill=S.PANEL, line=S.RULE, lw=0.75,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06,
             name=f"stage {i + 1}")
        dot(s, x + 0.1, y0 + 0.12, 0.34, S.TEAL, name="badge")
        text(s, x + 0.1, y0 + 0.12, 0.34, 0.34,
             [dict(runs=[(str(i + 1), F(15, bold=True, color=S.WHITE))],
                   align="c")], anchor="m")
        text(s, x + 0.5, y0 + 0.08, cw_ - 0.56, 0.62,
             [dict(runs=[(t_, F(15.5, bold=True))], ls=0.9)], anchor="m")
        text(s, x + 0.1, y0 + 0.74, cw_ - 0.2, 1.0,
             [dict(runs=[(e_, F(14, color=S.NAVY2))], ls=0.95)])
        pic(s, os.path.join(A, img), x + 0.08, y0 + 1.78, cw_ - 0.16)
        text(s, x + 0.06, y0 + ch - 0.42, cw_ - 0.12, 0.36,
             [dict(runs=[(cap, F(11.5, italic=True, color=S.GREY))],
                   align="c")], anchor="m")
        if i < 5:
            xa = x + cw_ + 0.005
            arrow(s, xa, y0 + 0.29, xa + gap - 0.01, y0 + 0.29,
                  color=S.TEAL, lw=2.0)
    by = y0 + ch + 0.12
    for (i0, i1, t_, col) in ((0, 3, "measured from the pixels", S.TEAL),
                              (4, 5, "statistics on those measurements",
                               S.NAVY2)):
        x0 = ML + i0 * (cw_ + gap)
        x1 = ML + i1 * (cw_ + gap) + cw_
        line(s, x0, by, x1, by, color=col, lw=1.25)
        line(s, x0, by - 0.07, x0, by, color=col, lw=1.25)
        line(s, x1, by - 0.07, x1, by, color=col, lw=1.25)
        text(s, x0, by + 0.02, x1 - x0, 0.3,
             [dict(runs=[(t_, F(13, italic=True, color=col))], align="c")])
    yb = 5.82
    rect(s, ML, yb, 5.55, 1.12, fill=S.TEAL_LIGHT, name="assumption box")
    text(s, ML + 0.15, yb + 0.04, 5.3, 1.05, [
        dict(runs=[("Assumption from step 2 on: ", F(14.5, bold=True)),
                   ("bright = silicon-based, grey = graphite. This is read "
                    "from brightness, not confirmed by chemical analysis "
                    "(e.g. EDS).", F(14.5))], ls=0.95)], anchor="m")
    xb = ML + 5.85
    bw = CW - 5.85
    rect(s, xb, yb, bw, 1.12, line=S.GREY, lw=1.25,
         dash=MSO_LINE_DASH_STYLE.DASH, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         radius=0.08, name="optional branch")
    pic(s, os.path.join(A, "thumb_dfn.png"), xb + bw - 1.62, yb + 0.16,
        h=0.8)
    text(s, xb + 0.15, yb + 0.03, bw - 1.85, 1.06, [
        dict(runs=[("Optional branch — a battery model",
                    F(14.5, bold=True))]),
        dict(runs=[("Image porosity, phase shares and particle sizes feed a "
                    "published cell model (PyBaMM DFN); everything else is "
                    "assumed.", F(13, color=S.NAVY2))], ls=0.92, sb=3)],
         anchor="m")
    x4 = ML + 3 * (cw_ + gap) + cw_ / 2
    line(s, x4, by + 0.33, x4, yb, color=S.GREY, lw=1.25,
         dash=MSO_LINE_DASH_STYLE.DASH, tail="triangle")
    text(s, x4 + 0.1, by + 0.31, 2.4, 0.26,
         [dict(runs=[("assumptions begin here", F(12, italic=True,
                                                  color=S.GREY))])],
         anchor="m")
    footer(s, 2, N)
    notes(s, f"""
WHAT TO SAY
Walk left to right. (1) Each location is photographed by three detectors. (2) Every pixel is labelled pore, grey bulk or bright particle with one fixed recipe; the brightness cut-offs are read from each photo's own histogram. (3) Quality guards check contrast, noise, sharpness and shading; {N.t('n_flagged')} photos fail ({N.t('flagged_ids')}, both Batch 1). (4) The label map becomes {N.t('n_features')} measurements; only those that agree between the left and right half of the same photo, and are not duplicates, are kept ({N.t('n_kept')}). (5) Batch 1 and Batch 2 are compared with Batch 3, asking whether a shift is bigger than normal photo-to-photo scatter. (6) A simple classifier says which batch a new sample is closest to. The dashed branch is a battery simulation that uses a few image numbers plus many assumptions. Steps 1-4 (teal bracket) are measurement; steps 5-6 add statistics on those measurements; in the dashed branch, assumptions dominate.

TERMS
- Detectors: BSE (composition contrast), Inlens (fine surface detail), ETD or SE (surface topography). Main uses Inlens only to refine fine cracks next to BSE pores, and texture features from Inlens/ETD.
- Quality guard: a fixed rule on an image property (e.g. bright-vs-grey contrast < 2.8 flags the photo).
- Repeatable: the value computed on the left half of a photo correlates with the right half (r >= 0.5 across photos).
- Centroid: the average of a batch's measurements; the classifier picks the nearest one.

HOW THE NUMBERS WERE OBTAINED (code -> step mapping, main @ {N.t('sha7')})
1. analysis/common.py find_locations(), load_gray() (median of RGB, crop coloured marker stripes), pixel_size_um() (TIFF tags).
2. analysis/segment.py segment(): normalise() 1st-99.5th percentile + Gaussian sigma 1.5 px; thresholds() = 3-class multi-Otsu (bright valley search falls back to multi-Otsu on {N.t('n_bright_fallback')}/31 photos); bright_particles() opening r=2, remove < 150 px, fill holes, watershed; pores_at() remove < 24 px; inlens_refine() adds Inlens-dark pixels within 2 px of a pore.
3. segment.py quality_guards() + GUARDS: contrast < 2.8, noise MAD > 11.2, sharpness > 11.0, shading > 55. Flagged photos stay in the data; results are reported with and without them.
4. features.py compute() ({N.t('n_features')} features), repeatability() (left/right halves segmented independently), prune(): drop if L/R r < 0.5, >50% missing/constant, or |r| > 0.9 with an already kept feature -> {N.t('n_kept')} kept.
5. baseline.py (Batch 3 median/MAD, 5-95% band) and deltas.py: mean shift, bootstrap 95% CI resampling photos, Mann-Whitney p, Holm correction over the {N.t('n_tests_per_batch')} kept features; classes: different (CI excludes 0 and Holm p < 0.05), suggestive (CI excludes 0 only), equivalent (CI within +/-10% of the Batch 3 mean), undetermined.
6. categorise.py: top-6 features by separation, robust Mahalanobis distance to the Batch 3 envelope (IN/OUT at its 95th percentile) and nearest batch centroid with softmax "surety"; validated leave-one-out ({N.t('loo_correct')}/31 correct; always-"Batch 3" = {N.t('majority_acc')}; shuffled labels {N.t('shuffled_acc')}).
Branch: dfn.py runs pybamm BasicDFNComposite with Chen2020_composite, setting porosity (measured, or 0.30), graphite/silicon volume shares and particle radii from the images; 30 resamples per batch at C/10, C/2, 1C, 2C. Electrode thickness, chemistry, Bruggeman exponent and all other parameters are assumed.

ASSUMPTIONS AND LIMITS
- Each photo is treated as one independent sample of its batch.
- Two documented reproduction issues on main (do not change the results shown): `run_all.py --clean` crashes because it deletes outputs/ after the folders were created; the DFN stage only re-runs when its output already exists (inverted check in run_all.stage_dfn).

SOURCES
- Pipeline driver: {repo_link(N, 'analysis/run_all.py')}
- Segmentation: {repo_link(N, 'analysis/segment.py', r'^def segment\(')}
- Feature gates: {repo_link(N, 'analysis/features.py', r'^def prune\(')}
- Comparison classes: {repo_link(N, 'analysis/deltas.py', r'def classify\(r\)')}
- Classifier: {repo_link(N, 'analysis/categorise.py')}
- Simulation: {repo_link(N, 'analysis/dfn.py')}; {DOI['chen'][0]} {DOI['chen'][1]}; {DOI['pybamm'][0]} {DOI['pybamm'][1]}

LIKELY QUESTION
Q: "Is this machine learning?"
A: No black box. Every step is a fixed, inspectable rule (thresholds, morphology, distances). The only fitted parts are averages and spreads of the measurements, and the classifier is a nearest-average rule. That makes it explainable, but it also means it can only be as good as the brightness labels and the small number of photos.{rebuild_caveat(N)}
{provenance_block(N)}""")
    return s


def slide3(prs, N, L, ann):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    a = ann["seg"]
    title(s, "How we segment the images",
          "One recipe for every photo; the cut-offs come from each photo's "
          "own brightness histogram")
    pw, gap = (CW - 2 * 0.2) / 3, 0.2
    ph = pw * a["h"] / a["w"]
    y = 1.98
    rw, rh = a["region"][2], a["region"][3]
    labels = ["1 · Original photo (BSE)",
              "2 · Brightness cut-offs, before clean-up",
              "3 · Final labels"]
    imgs = ["seg_original.png", "seg_threshold.png", "seg_overlay.png"]
    for i in range(3):
        x = ML + i * (pw + gap)
        text(s, x, y - 0.36, pw, 0.32,
             [dict(runs=[(labels[i], F(14.5, bold=True))])], anchor="b")
        pic(s, os.path.join(A, imgs[i]), x, y, pw, ph, name=labels[i])
    fx = lambda i, px: ML + i * (pw + gap) + pw * px / rw    # noqa: E731
    fy = lambda py: y + ph * py / rh                          # noqa: E731
    scalebar(s, ML + 0.2, y + ph - 0.14, pw * 400 / rw, "10 μm", size=11)
    label(s, ML + pw - 2.05, y + ph - 0.38, 1.95, 0.28,
          f"{N.t('ex_batch')} · img_{N.t('ex_id')}", size=11, color=S.WHITE,
          fill=S.NAVY, alpha=0.72, align="c")
    cx0, cy0, cw0, ch0 = a["closeup"]
    bx, byy = fx(2, cx0 - a["region"][0]), fy(cy0 - a["region"][1])
    rect(s, bx, byy, pw * cw0 / rw, ph * ch0 / rh, line=S.TEAL, lw=2.0,
         dash=MSO_LINE_DASH_STYLE.DASH, name="close-up box")
    x3 = ML + 2 * (pw + gap)
    placed = []
    for (px, py), t_ in ((a["particle"], "bright particle"),
                         (a["split"], "split line"), (a["pore"], "pore")):
        tx, ty = fx(2, px), fy(py)
        lw_, lh_ = tw(t_, 12.5) + 0.26, 0.3
        best = None
        for dx, dy in ((0.25, -0.62), (-0.25 - lw_, -0.62), (0.25, 0.32),
                       (-0.25 - lw_, 0.32), (0.45, -0.15), (-0.45 - lw_, -0.15)):
            lx, ly = tx + dx, ty + dy
            if lx < x3 + 0.05 or lx + lw_ > x3 + pw - 0.05 or \
                    ly < y + 0.05 or ly + lh_ > y + ph - 0.05:
                continue
            if any(not (lx + lw_ < qx or qx + qw_ < lx or ly + lh_ < qy or
                        qy + 0.3 < ly) for qx, qy, qw_ in placed):
                continue
            best = (lx, ly)
            break
        lx, ly = best or (min(max(tx - lw_ / 2, x3 + 0.05),
                              x3 + pw - lw_ - 0.05), y + 0.08)
        placed.append((lx, ly, lw_))
        ax_ = min(max(tx, lx + 0.12), lx + lw_ - 0.12)
        ay_ = ly + lh_ if ly < ty else ly
        arrow(s, ax_, ay_, tx, ty, color=S.NAVY, lw=1.5)
        label(s, lx, ly, lw_, lh_, t_, size=12.5)
    ly = y + ph + 0.1
    legend = [(S.PORE_RGB, "pore / crack (dark)"),
              (S.BRIGHT_RGB, "bright particle — assumed silicon-based"),
              ((150, 150, 150), "grey bulk — assumed graphite"),
              (None, "particle outline (watershed)")]
    lx = ML
    for col, t_ in legend:
        if col is None:
            rect(s, lx, ly + 0.07, 0.26, 0.2, fill="%02X%02X%02X" % S.BRIGHT_RGB, line="232323",
                 lw=1.5)
        else:
            rect(s, lx, ly + 0.07, 0.26, 0.2, fill="%02X%02X%02X" % col,
                 line=S.NAVY2, lw=0.5)
        text(s, lx + 0.33, ly, tw(t_, 13.5) + 0.3, 0.34,
             [dict(runs=[(t_, F(13.5))])], anchor="m")
        lx += 0.33 + tw(t_, 13.5) + 0.5
    L.add(3, "Panels show Batch 3 photo cfe5vt7s", "ex_id", "ex_batch")

    yb = ly + 0.42
    steps = [
        ("Stretch & smooth", "1st–99.5th percentile, 1.5-px blur"),
        ("Two cut-offs", "darkest group → pore, brightest → particle"),
        ("Clean up", "drop specks, fill holes, add Inlens cracks"),
        ("Split touching particles", "watershed on distance to edge"),
    ]
    paras = [dict(runs=[("What the recipe does", F(14.5, bold=True,
                                                   color=S.TEAL))])]
    for i, (h_, b_) in enumerate(steps):
        paras.append(dict(runs=[(f"{i + 1}  {h_}: ", F(12.5, bold=True)),
                                (b_, F(12.5, color=S.NAVY2))], ls=0.9, sb=3))
    paras.append(dict(runs=[("Why: ", F(12.5, bold=True, color=S.TEAL)),
                            ("transparent, no training, same settings for "
                             "all photos", F(12.5, color=S.NAVY2))], ls=0.9,
                      sb=5))
    text(s, ML, yb, 5.45, 1.8, paras)
    hx = ML + 5.62
    text(s, hx, yb, 3.0, 0.3,
         [dict(runs=[("This photo's histogram", F(13, bold=True))],
               align="c")])
    pic(s, os.path.join(A, "seg_histogram.png"), hx + 0.25, yb + 0.31, 2.5)
    text(s, hx, yb + 1.5, 3.0, 0.28, [
        dict(runs=[(f"pore cut {N.t('ex_t_lo')}", F(12, color="B5172F")),
                   ("  ·  ", F(12, color=S.GREY)),
                   (f"bright cut {N.t('ex_t_b')}", F(12, color="8A6200"))],
             align="c")])
    L.add(3, "Cut-offs for this photo: pore 0.277, bright 0.607 "
          "(normalised scale)", "ex_t_lo", "ex_t_b")
    cxp = hx + 3.2
    cwid = ML + CW - cxp
    ch_w = (cwid - 0.08) / 2
    ch_h = ch_w * ch0 / cw0
    if ch_h > 1.0:
        ch_h = 1.0
        ch_w = ch_h * cw0 / ch0
        cwid = 2 * ch_w + 0.08
    text(s, cxp, yb, cwid, 0.3,
         [dict(runs=[("Close-up: an ambiguous object", F(13, bold=True))])])
    pic(s, os.path.join(A, "closeup_original.png"), cxp, yb + 0.31, ch_w,
        ch_h)
    pic(s, os.path.join(A, "closeup_overlay.png"), cxp + ch_w + 0.08,
        yb + 0.31, ch_w, ch_h)
    rect(s, cxp, yb + 0.31, cwid, ch_h, line=S.TEAL, lw=1.5,
         dash=MSO_LINE_DASH_STYLE.DASH)
    text(s, cxp, yb + 0.35 + ch_h, ML + CW - cxp, 0.45, [dict(runs=[(
        "Textured light-grey object: only partly labelled bright, and "
        f"split into {a['closeup_pieces']} ‘particles’.",
        F(12, color=S.NAVY2))], ls=0.9)])
    rect(s, ML, 6.46, CW, 0.56, fill=S.PANEL)
    text(s, ML + 0.15, 6.46, CW - 0.3, 0.56, [dict(runs=[
        ("Main limitation: ", F(14, bold=True, color=S.AMBER)),
        ("brightness is not chemistry — ‘bright = silicon’ is unconfirmed, "
         "low-contrast photos inflate the bright class, and gaps under ~50 nm "
         "(2 px) are invisible.", F(14))])], anchor="m")
    footer(s, 3, N)
    L.add(3, "+/-10% shift of the pore cut changes pore fraction by -12%/+13% "
          "(median, relative)", "sens_pore_m", "sens_pore_p")
    L.add(3, "+/-10% shift of the bright cut changes bright fraction by "
          "+23%/-10%", "sens_si_m", "sens_si_p")
    L.add(3, "Human point-count check sheet: 300 points, 0 labelled",
          "pc_points", "pc_human")
    L.add(3, "Bright cut falls back to multi-Otsu on all 31 photos",
          "n_bright_fallback")
    notes(s, f"""
WHAT TO SAY
Same patch of a Batch 3 photo, three times. Left: the photo as the BSE detector sees it. Middle: the two brightness cut-offs applied directly — red is darker than the pore cut, gold brighter than the bright cut. You can see why clean-up is needed: thin bright rims on flake edges and tiny specks would be counted as particles. Right: the final labels after removing specks, filling holes and splitting touching particles (dark outlines). The close-up shows an honest ambiguous case: a textured light-grey object is only partly called bright and gets cut into {a['closeup_pieces']} pieces. The main limitation is that brightness is not chemistry.

TERMS
- Histogram: how many pixels have each brightness.
- Multi-Otsu: an automatic rule that splits the histogram into three groups so each group is as uniform as possible (Otsu 1979; 3-class version in scikit-image).
- Watershed: treats each particle like a landscape (distance to its edge) and draws split lines along the valleys between touching particles.
- Inlens refinement: the Inlens detector sees fine cracks better; dark Inlens pixels within 2 px of a BSE pore are added to the pore class.

HOW THE NUMBERS WERE OBTAINED
- Panels: Batch_3/img_{N.t('ex_id')}_BSE.tif, region x {a['region'][0]}-{a['region'][0] + a['region'][2]}, y {a['region'][1]}-{a['region'][1] + a['region'][3]} px ({a['region'][2] * 0.025:g} x {a['region'][3] * 0.025:g} um). Panel 2 = main's normalise() output with this photo's committed cut-offs (segment_info.csv: t_lo = {N.t('ex_t_lo')}, t_bright = {N.t('ex_t_b')}); panel 3 = main's committed masks (analysis/outputs/masks/Batch_3_{N.t('ex_id')}_masks.npz), which src/verify_main.py reproduced pixel-for-pixel.
- Per-photo cut-offs: the bright "valley" search never finds a separate bright peak, so the upper multi-Otsu cut is used on {N.t('n_bright_fallback')}/31 photos (segment_info.csv, bright_method).
- Robustness (threshold_sensitivity.csv, median over photos): shifting the pore cut by -10%/+10% changes the pore fraction by {N.t('sens_pore_m')}/{N.t('sens_pore_p')}; shifting the bright cut by -10%/+10% changes the bright fraction by {N.t('sens_si_m')}/{N.t('sens_si_p')}. Repeatability (left vs right half, repeatability.csv): pore fraction r = {N.t('lr_pore_frac')}, bright fraction r = {N.t('lr_silicon_frac')}.
- Human check: analysis/outputs/pointcount/point_count_validation.xlsx has {N.t('pc_points')} random points with crops, but {N.t('pc_human')} human labels — segmentation accuracy against an expert is therefore UNAVAILABLE.

ASSUMPTIONS AND LIMITS
- Bright = silicon-based and grey = graphite are assumptions from appearance (main's standing assumption). There is no separate class for binder, coatings or "other", so anything bright enough is counted as particle.
- Because cut-offs adapt to each photo, a photo with weak contrast pushes grey material into the bright class: the two flagged Batch 1 photos read {N.t('si_flagged_values')} bright versus {N.t('si_range_unflagged')} for the other 29.
- Pixels are 25 nm; pores narrower than ~2 px (50 nm) and objects below the size filters are not measured, so visible porosity (~10%) understates true porosity.
- Watershed can over-split irregular particles (close-up), which lowers sizes and raises counts.

SOURCES
- {repo_link(N, 'analysis/segment.py', r'^def segment\(')}
- Settings: {repo_link(N, 'analysis/common.py', r'^P_LO, P_HI')}
- {DOI['otsu'][0]} {DOI['otsu'][1]}

LIKELY QUESTION
Q: "How do you know the labels are right?"
A: We know they are stable — re-running main's code reproduces every mask exactly, left and right halves agree, and +/-10% threshold changes move fractions by about 10-25% in relative terms. We do not yet know they are correct: the 300-point expert check sheet in the repo is still empty, and nobody has confirmed with chemical analysis that bright means silicon.
{provenance_block(N)}""")
    return s


def slide4(prs, N, L):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "The markers that matter — and what we can conclude")
    text(s, ML, 1.13, CW, 0.42, [dict(runs=[(
        f"Five of the {N.t('n_kept')} retained measurements, chosen for "
        "reliability and physical relevance", F(17, color=S.NAVY2))])])
    feat = "analysis/features.py"
    ref_num = {"singh": 1, "obrovac": 2, "liu12": 3, "wucui": 4, "yolk": 5,
               "ebner": 6, "billaud": 7}

    def refs(*keys):
        out = [(" [", F(12.5, color=S.TEAL))]
        for i, k in enumerate(keys):
            if i:
                out.append((",", F(12.5, color=S.TEAL)))
            out.append((str(ref_num[k]), F(12.5, color=S.TEAL,
                                           link=DOI[k][1])))
        out.append(("]", F(12.5, color=S.TEAL)))
        return out

    def marker(name, var):
        return [dict(runs=[(name, F(13.5, bold=True))], ls=0.92),
                dict(runs=[(var, F(11, font=S.MONO_FONT, color=S.TEAL,
                                   link=repo_link(N, feat, rf'^F\("{var}"')))],
                     sb=2)]

    def body(t_, extra=None):
        runs = [(t_, F(13, color=S.NAVY2))]
        if extra:
            runs += extra
        return [dict(runs=runs, ls=0.92)]

    def concl(tag, col, t_):
        return [dict(runs=[(tag + " ", F(13, bold=True, color=col)),
                           (t_, F(13))], ls=0.92)]

    d = lambda b, f: N.pct(f"d_{b}_{f}_pct", signed=True)       # noqa: E731
    dn = lambda b, f: N.pct(f"d_nf_{b}_{f}_pct", signed=True)   # noqa: E731
    rows = [
        (marker("Visible pore space", "pore_frac"),
         body("Share of the section that is dark: gaps wider than ~50 nm"),
         body("Pores carry the electrolyte: more pore space speeds ion "
              "transport but stores less energy per volume", refs("singh")),
         concl("Tentative.", S.AMBER,
               f"Batch 1 is {d('Batch_1', 'pore_frac')} vs Batch 3 (interval "
               "excludes zero, not significant after correction). Batch 2: "
               "no detected difference.")),
        (marker("Bright-particle share", "silicon_frac"),
         body("Share of bright particles — assumed silicon-based, not "
              "chemically confirmed"),
         body("Silicon holds ~10× more lithium per gram than graphite but "
              "swells up to ~280%", refs("obrovac")),
         concl("No difference detected", S.TEAL,
               f"once two low-contrast Batch 1 photos are set aside "
               f"({dn('Batch_1', 'silicon_frac')}); with them Batch 1 reads "
               f"{d('Batch_1', 'silicon_frac')}.")),
        (marker("Bright-particle size", "si_d50_um"),
         body("Median diameter of the bright particles in a photo "
              "(equal-area circle)"),
         body("Smaller particles crack less when they swell", refs("liu12")
              + [(" but expose more surface to side reactions",
                  F(13, color=S.NAVY2))] + refs("wucui")),
         concl("Undetermined:", S.GREY,
               f"Batch 3 averages {N.v('mean_si_d50_um_Batch_3'):.2f} μm; "
               f"Batch 1 {d('Batch_1', 'si_d50_um')}, Batch 2 "
               f"{d('Batch_2', 'si_d50_um')}, both within the noise. A 2-D "
               "cut understates true size.")),
        (marker("Pore space next to particles", "ring_porosity_250nm"),
         body("Pore share of a 250-nm shell around each bright particle"),
         body("Nearby voids can absorb swelling and let electrolyte in",
              refs("yolk") + [("; too much may weaken contact",
                               F(13, color=S.NAVY2))]),
         concl("No difference detected:", S.TEAL,
               f"{100 * min(N.v('mean_ring_porosity_250nm_Batch_2'), N.v('mean_ring_porosity_250nm_Batch_3')):.0f}–"
               f"{100 * N.v('mean_ring_porosity_250nm_Batch_1'):.0f}% on "
               "average in every batch.")),
        (marker("Graphite flake alignment", "gr_chord_ratio_hv"),
         body("Horizontal vs vertical run length through the grey flakes"),
         body("Flat-lying flakes lengthen the ion path through the coating "
              "thickness", refs("ebner", "billaud")),
         concl("Shared feature:", S.TEAL,
               f"about {N.v('mean_gr_chord_ratio_hv_Batch_3'):.1f} in every batch "
               "(flakes ~20% wider than tall); equivalent within ±10%.")),
    ]
    L.add(4, "Batch 1 visible pore space -14% vs Batch 3, suggestive (CI "
          "excludes 0, Holm p = 1)", "d_Batch_1_pore_frac_pct",
          "d_Batch_1_pore_frac_class", "d_Batch_1_pore_frac_ci",
          "d_Batch_1_pore_frac_pholm")
    L.add(4, "Batch 2 visible pore space: undetermined",
          "d_Batch_2_pore_frac_class")
    L.add(4, "Bright share: Batch 1 -2% without flagged photos, +45% with "
          "them (both undetermined)", "d_nf_Batch_1_silicon_frac_pct",
          "d_Batch_1_silicon_frac_pct", "d_Batch_1_silicon_frac_class",
          "d_nf_Batch_1_silicon_frac_class")
    L.add(4, "Particle size: Batch 3 mean 0.76 um; B1 -15%, B2 -12%, "
          "undetermined", "mean_si_d50_um_Batch_3", "d_Batch_1_si_d50_um_pct",
          "d_Batch_2_si_d50_um_pct", "d_Batch_1_si_d50_um_class",
          "d_Batch_2_si_d50_um_class")
    L.add(4, "Pore space within 250 nm of particles: 8-10% in every batch, "
          "no difference detected", "mean_ring_porosity_250nm_Batch_1",
          "mean_ring_porosity_250nm_Batch_2",
          "mean_ring_porosity_250nm_Batch_3",
          "d_Batch_1_ring_porosity_250nm_class",
          "d_Batch_2_ring_porosity_250nm_class")
    L.add(4, "Flake alignment ratio ~1.2 in every batch, equivalent",
          "mean_gr_chord_ratio_hv_Batch_1", "mean_gr_chord_ratio_hv_Batch_2",
          "mean_gr_chord_ratio_hv_Batch_3", "d_Batch_1_gr_chord_ratio_hv_class",
          "d_Batch_2_gr_chord_ratio_hv_class")
    booktabs(s, ML, 1.62, [2.45, 2.85, 3.45, 3.483],
             [0.42] + [0.84] * 5,
             ["Marker (plain English)", "What the image measurement tells us",
              "Why it might matter for a battery",
              "What we can actually conclude"], rows, name="marker table")
    rl = []
    for k, i in ref_num.items():
        au = DOI[k][0].split(" (")[0].replace(", Kaiser & Hahn", " et al.") \
            .replace(", Chung, Garcia & Wood", " et al.")
        yr = DOI[k][0].split("(")[1][:4]
        rl += [(f"[{i}] ", F(11, color=S.TEAL)),
               (f"{au} {yr}", F(11, color=S.NAVY2, link=DOI[k][1])),
               ("   ", F(11))]
    text(s, ML, 6.42, CW, 0.62, [
        dict(runs=[("Comparisons: mean shift vs Batch 3 with a bootstrap 95% "
                    f"interval, Mann–Whitney test, Holm correction over "
                    f"{N.t('n_tests_per_batch')} markers; "
                    f"{N.t('n_Batch_1')} vs {N.t('n_Batch_3')} photos. "
                    "Variable names link to their definitions in main.",
                    F(11, italic=True, color=S.GREY))]),
        dict(runs=rl, sb=2)])
    footer(s, 4, N)
    notes(s, f"""
WHAT TO SAY
These five are the markers I would show anyone first. They were picked because they are repeatable on main, have a clear physical meaning, and together describe composition, size, contact and arrangement — not because they show the biggest batch differences. Read the last column carefully: only one is even tentatively different (Batch 1's lower visible pore space), two show no detected difference, one is undetermined, and one is a solid shared feature. Every battery effect in the third column is conditional: "might matter", with trade-offs in both directions.

TERMS
- Pore / porosity: empty space between solids, normally filled with liquid electrolyte in a cell.
- Equal-area diameter: the diameter of a circle with the same area as the particle's cut.
- Run length (chord): length of an uninterrupted straight line through one phase.
- Interval (bootstrap 95%): range of plausible mean shifts from re-sampling the photos.
- Holm correction: makes the significance test stricter because 20 markers are tested per batch.
- Tentative / suggestive: interval excludes zero but the corrected test is not significant. Undetermined: interval includes zero and is too wide to call the batches equivalent. Equivalent: interval within +/-10% of the Batch 3 mean.

HOW THE NUMBERS WERE OBTAINED (exact variables, main @ {N.t('sha7')})
- pore_frac (features.py): pore pixels / all pixels; Batch 3 mean {100 * N.v('mean_pore_frac_Batch_3'):.1f}%; Batch 1 {d('Batch_1', 'pore_frac')} (deltas.csv: CI {N.t('d_Batch_1_pore_frac_ci')}, Mann-Whitney p = {N.t('d_Batch_1_pore_frac_pmw')}, Holm p = {N.t('d_Batch_1_pore_frac_pholm')}, class {N.t('d_Batch_1_pore_frac_class')}); Batch 2 {d('Batch_2', 'pore_frac')} ({N.t('d_Batch_2_pore_frac_class')}). Repeatability r = {N.t('lr_pore_frac')}.
- silicon_frac: bright pixels / all pixels. Batch 1 {d('Batch_1', 'silicon_frac')} with all 7 photos ({N.t('d_Batch_1_silicon_frac_class')}), {dn('Batch_1', 'silicon_frac')} without the 2 flagged photos (deltas_noflag.csv, {N.t('d_nf_Batch_1_silicon_frac_class')}). r = {N.t('lr_silicon_frac')}.
- si_d50_um: median (number-weighted) equal-area diameter of watershed particles per photo; Batch 3 mean {N.v('mean_si_d50_um_Batch_3'):.3f} um; Batch 1 {d('Batch_1', 'si_d50_um')}, Batch 2 {d('Batch_2', 'si_d50_um')} (both {N.t('d_Batch_1_si_d50_um_class')}). r = {N.t('lr_si_d50_um')}. Main removed si_d10_um because 2-D cuts through packed particles create a ~0.37 um floor.
- ring_porosity_250nm: pore fraction of the 10-px (~250 nm) shell around all bright particles; batch means {100 * N.v('mean_ring_porosity_250nm_Batch_1'):.1f}% / {100 * N.v('mean_ring_porosity_250nm_Batch_2'):.1f}% / {100 * N.v('mean_ring_porosity_250nm_Batch_3'):.1f}% (B1/B2/B3); classes {N.t('d_Batch_1_ring_porosity_250nm_class')} / {N.t('d_Batch_2_ring_porosity_250nm_class')}. r = {N.t('lr_ring_porosity_250nm')}.
- gr_chord_ratio_hv: mean horizontal / mean vertical run length through the grey phase; means {N.v('mean_gr_chord_ratio_hv_Batch_1'):.3f} / {N.v('mean_gr_chord_ratio_hv_Batch_2'):.3f} / {N.v('mean_gr_chord_ratio_hv_Batch_3'):.3f}; both comparisons {N.t('d_Batch_1_gr_chord_ratio_hv_class')}. r = {N.t('lr_gr_chord_ratio_hv')}.
- Not prioritised but available: pores_per_mm2 (pore count density; Batch 1 {d('Batch_1', 'pores_per_mm2')}, {N.t('d_Batch_1_pores_per_mm2_class')}; Batch 2 {d('Batch_2', 'pores_per_mm2')}, {N.t('d_Batch_2_pores_per_mm2_class')}). Counts depend on resolution and on the 24-px minimum pore size, so compare them only between photos taken with the same settings. Pore orientation (pore_anisotropy) was dropped by main's own repeatability gate (r = {N.t('lr_pore_anisotropy')} < 0.5).

ASSUMPTIONS AND LIMITS
- With 7 photos per batch, real differences of about one Batch 3 standard deviation can sit in "undetermined" (main lists the minimum detectable difference per marker in deltas.csv, column mdd).
- Everything is a 2-D section: sizes, contact and alignment are section statistics, not 3-D measurements.
- Battery effects are literature mechanisms, not measured on these electrodes.

SOURCES (full references; the bracket numbers on the slide are clickable DOIs)
[1] {DOI['singh'][0]} {DOI['singh'][1]}
[2] {DOI['obrovac'][0]} {DOI['obrovac'][1]}
[3] {DOI['liu12'][0]} {DOI['liu12'][1]} (crystalline Si particles below ~150 nm did not fracture on first lithiation)
[4] {DOI['wucui'][0]} {DOI['wucui'][1]}
[5] {DOI['yolk'][0]} {DOI['yolk'][1]} (engineered void space lets Si expand)
[6] {DOI['ebner'][0]} {DOI['ebner'][1]}
[7] {DOI['billaud'][0]} {DOI['billaud'][1]}
Definitions: {repo_link(N, feat)} ; comparison rule: {repo_link(N, 'analysis/deltas.py', r'def classify\(r\)')}

LIKELY QUESTION
Q: "Is more pore space good or bad?"
A: Neither by itself. More pore space usually helps ions move (better fast charging and discharging) but leaves less room for active material, so the cell stores less energy per volume. The right amount depends on the cell design — which is why we report the shift and its uncertainty, not a verdict.
{provenance_block(N)}""")
    return s


def slide5(prs, N, L, ann):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "What these markers look like",
          f"One typical {N.t('ex_batch')} photo (img_{N.t('ex_id')}) · all "
          "outlines are algorithm labels, none expert-confirmed")
    E = ann["examples"]
    qw = (CW - 0.3) / 2
    pos = [(ML, 1.72), (ML + qw + 0.3, 1.72), (ML, 4.14),
           (ML + qw + 0.3, 4.14)]
    iw = 3.05

    def frame(i, key, img, head, kind="main metric"):
        x, y = pos[i]
        e = E[key]
        ih = iw * e["h"] / e["w"]
        pic(s, os.path.join(A, img), x, y, iw, ih, name=head)
        tx = x + iw + 0.2
        tw_ = qw - iw - 0.2
        text(s, tx, y - 0.06, tw_, 0.34,
             [dict(runs=[(head, F(16, bold=True))])])
        text(s, tx, y + 0.29, tw_, 0.22, [dict(runs=[(kind.upper(), F(
            9.5, bold=True, color=S.TEAL if kind == "main metric" else
            S.AMBER))])])
        fxx = lambda px: x + iw * px / e["w"]       # noqa: E731
        fyy = lambda py: y + ih * py / e["h"]       # noqa: E731
        return tx, tw_, fxx, fyy, ih

    # (a) size ---------------------------------------------------------------
    e = E["size"]
    tx, tw_, fx, fy, ih = frame(0, "size", "ex_size.png", "Particle size")
    pts = [(fx(px), fy(py)) for px, py in e["contour"]]
    freeform(s, pts, color=S.TEAL, lw=2.5, name="particle outline")
    r_in = (iw * e["eq_diam_px"] / e["w"]) / 2
    cx, cy = fx(e["centroid"][0]), fy(e["centroid"][1])
    rect(s, cx - r_in, cy - r_in, 2 * r_in, 2 * r_in, line=S.WHITE, lw=1.5,
         dash=MSO_LINE_DASH_STYLE.DASH, shape=MSO_SHAPE.OVAL,
         name="equal-area circle")
    line(s, cx - r_in, cy, cx + r_in, cy, color=S.WHITE, lw=2.0,
         head="triangle", tail="triangle", name="diameter")
    dl = f"d = {e['eq_diam_um']:.2f} μm"
    lw_ = tw(dl, 12, True) * 1.35 + 0.3
    label(s, cx - lw_ / 2, cy + 0.08, lw_, 0.29, dl, size=12, bold=True,
          align="c")
    scalebar(s, pos[0][0] + 0.14, pos[0][1] + ih - 0.12, iw * 80 / e["w"],
             "2 μm", size=10)
    text(s, tx, pos[0][1] + 0.52, tw_, 1.6, [
        dict(runs=[(dl, F(20, bold=True, color=S.TEAL))]),
        dict(runs=[("diameter of a circle with the same area as this "
                    "particle's cut", F(12.5, color=S.NAVY2))], ls=0.9),
        dict(runs=[(f"Photo median over {N.t('ex_n_particles')} particles: "
                    f"{N.v('ex_si_d50_um'):.2f} μm — many small ones pull it "
                    "down.", F(13))], ls=0.9, sb=6)])
    L.add(5, f"Highlighted particle: equal-area diameter "
          f"{e['eq_diam_um']:.2f} um (main particles.csv, label "
          f"{e['table_label']})", note="value read from main's particles.csv "
          "(diam_um); verified against the committed label mask")
    L.add(5, "Photo median particle size over 327 particles",
          "ex_si_d50_um", "ex_n_particles")

    # (b) pores --------------------------------------------------------------
    e = E["pores"]
    tx, tw_, fx, fy, ih = frame(1, "pores", "ex_pores.png",
                                "Visible pore space")
    px, py = fx(e["target"][0]), fy(e["target"][1])
    label(s, px + 0.35, py + 0.5, 0.95, 0.3, "pore", size=12.5)
    arrow(s, px + 0.6, py + 0.5, px, py, color=S.WHITE, lw=2.0)
    scalebar(s, pos[1][0] + 0.14, pos[1][1] + ih - 0.12, iw * 200 / e["w"],
             "5 μm", size=10)
    text(s, tx, pos[1][1] + 0.52, tw_, 1.6, [
        dict(runs=[(N.pct("ex_pore_frac", 1), F(20, bold=True,
                                                color=S.TEAL))]),
        dict(runs=[("of this photo's area is pore", F(12.5,
                                                      color=S.NAVY2))]),
        dict(runs=[("Red = darker than the pore cut-off: gaps wider than "
                    "~50 nm. Finer pores are invisible.", F(13))], ls=0.9,
             sb=6)])
    L.add(5, "Example photo pore fraction 10.8%", "ex_pore_frac")

    # (c) contact -------------------------------------------------------------
    e = E["contact"]
    tx, tw_, fx, fy, ih = frame(2, "contact", "ex_contact.png",
                                "Particle–pore contact",
                                kind="2 main metrics + 2 derived")
    px, py = fx(e["target"][0]), fy(e["target"][1])
    lt = "edge within 75 nm of a pore"
    lw_ = tw(lt, 11.5) + 0.25
    lx = min(max(px - lw_ / 2, pos[2][0] + 0.06), pos[2][0] + iw - lw_ - 0.06)
    label(s, lx, pos[2][1] + 0.08, lw_, 0.28, lt, size=11.5)
    arrow(s, min(max(px, lx + 0.2), lx + lw_ - 0.2), pos[2][1] + 0.36, px,
          py, color=S.WHITE, lw=2.0)
    scalebar(s, pos[2][0] + 0.14, pos[2][1] + ih - 0.12, iw * 200 / e["w"],
             "5 μm", size=10)
    rows = [
        (N.pct("ex_contact_number"), "of particles touch a pore", "derived"),
        (N.pct("ex_contact_area"), "of bright area is in them", "derived"),
        (f"{100 * N.v('ex_si_border_pore'):.1f}%", "of particle edge touches "
         "a pore", "main"),
        (N.pct("ex_ring_porosity_250nm", 1), "of a 250-nm shell is pore",
         "main"),
    ]
    paras = []
    for i, (v_, a_, b_) in enumerate(rows):
        paras.append(dict(runs=[(v_ + " ", F(15, bold=True, color=S.TEAL)),
                                (a_ + " ", F(12)),
                                (f"({b_})", F(10.5, italic=True,
                                              color=S.AMBER if b_ ==
                                              "derived" else S.GREY))],
                          ls=0.86, sb=2 if i else 0))
    text(s, tx, pos[2][1] + 0.5, tw_, 1.75, paras)
    L.add(5, "Share of particles with an edge pixel within 3 px (75 nm) of a "
          "pore (presentation-derived)", "ex_contact_number")
    L.add(5, "Share of bright area belonging to those particles "
          "(presentation-derived)", "ex_contact_area")
    L.add(5, "Share of particle-edge pixels within 3 px of a pore "
          "(main si_border_pore)", "ex_si_border_pore")
    L.add(5, "Pore share of the 250-nm shell around particles "
          "(main ring_porosity_250nm)", "ex_ring_porosity_250nm")

    # (d) alignment -----------------------------------------------------------
    e = E["align"]
    tx, tw_, fx, fy, ih = frame(3, "align", "ex_align.png",
                                "Flake alignment")
    for x1, y1, x2, y2 in e["hseg"]:
        line(s, fx(x1), fy(y1), fx(x2), fy(y2), color=S.TEAL_RUN, lw=2.5,
             head="oval", tail="oval", name="horizontal run")
    for x1, y1, x2, y2 in e["vseg"]:
        line(s, fx(x1), fy(y1), fx(x2), fy(y2), color=S.VRUN, lw=2.5,
             head="oval", tail="oval", name="vertical run")
    scalebar(s, pos[3][0] + 0.14, pos[3][1] + ih - 0.12, iw * 200 / e["w"],
             "5 μm", size=10)
    text(s, tx, pos[3][1] + 0.52, tw_, 1.7, [
        dict(runs=[(f"ratio {N.v('ex_gr_chord_ratio_hv'):.2f}",
                    F(20, bold=True, color=S.TEAL))]),
        dict(runs=[("mean run ", F(12.5, color=S.NAVY2)),
                   ("across", F(12.5, bold=True, color=S.TEAL_RUN)),
                   (f" {N.v('ex_gr_chord_h_um'):.2f} μm vs ",
                    F(12.5, color=S.NAVY2)),
                   ("down", F(12.5, bold=True, color=S.VRUN)),
                   (f" {N.v('ex_gr_chord_v_um'):.2f} μm",
                    F(12.5, color=S.NAVY2))], ls=0.9),
        dict(runs=[("Runs through the grey flakes are longer sideways: the "
                    "flakes lie flat.", F(13))], ls=0.9, sb=6)])
    L.add(5, "Example photo flake run ratio (across vs down)",
          "ex_gr_chord_ratio_hv", "ex_gr_chord_h_um", "ex_gr_chord_v_um")
    text(s, ML, 6.42, CW, 0.58, [dict(runs=[
        ("‘Contact’ here is a 2-D distance in one section ", F(13, bold=True,
                                                            color=S.NAVY2)),
        ("(edge within 75 nm of a pore) — not wetting, electrical contact or "
         "electrochemical activity. Main = a metric computed by main; derived "
         "= computed for this slide from main's per-particle table.",
         F(13, color=S.NAVY2))], ls=0.92)])
    footer(s, 5, N)
    notes(s, f"""
WHAT TO SAY
Here is what four of the markers literally measure, on one ordinary Batch 3 photo. Top left: one particle outlined, with the circle of equal area — its diameter is the size number; the photo's median is smaller because there are many tiny particles. Top right: everything red is "visible pore"; {N.pct('ex_pore_frac', 1)} of this photo. Bottom left: contact. I show four different numbers from the same photo on purpose — share of particles that touch a pore, share of area in those particles, share of particle edge next to a pore, and the pore share of a thin shell around particles. They answer different questions, so we should never say "contact" without saying which one. Bottom right: lines through the grey flakes are longer sideways than up-and-down, so the flakes lie flat. All of this is algorithm output; nobody has hand-checked these outlines yet.

TERMS
- Equal-area diameter: diameter of a circle whose area equals the particle's cut area.
- Shell (ring): the band within 10 pixels (~250 nm) outside every bright particle.
- Contact (here): a particle edge pixel within 3 pixels (75 nm) of a pore pixel in this 2-D section. It is not proof of electrolyte wetting, of an electrical connection, or of electrochemical activity.
- Run length: an uninterrupted straight line through the grey phase, measured row by row (across) or column by column (down).

HOW THE NUMBERS WERE OBTAINED
- Particle: watershed label {E['size']['table_label']} of img_{N.t('ex_id')}; diam_um = {E['size']['eq_diam_um']:.3f} um in analysis/outputs/tables/particles.csv; outline traced from the committed label mask. Photo median si_d50_um = {N.v('ex_si_d50_um'):.4f} um over {N.t('ex_n_particles')} particles (features.csv).
- Pore space: pore_frac = {N.v('ex_pore_frac'):.4f} (features.csv). The crop shown has {100 * E['pores']['crop_pore_frac']:.1f}% pore.
- Contact, main metrics: si_border_pore = {N.v('ex_si_border_pore'):.4f} (share of bright-mask edge pixels within 3 px of a pore; exp_newfeat.csv — passed main's gates but is not in the 20-marker decision set); ring_porosity_250nm = {N.v('ex_ring_porosity_250nm'):.4f} (features.csv, a retained marker). Contact, illustration only: {N.t('ex_contact_number')} of particles and {N.t('ex_contact_area')} of particle area have min distance to a pore <= 3 px (column d_min_px of main's particles.csv) — these two shares are derived for this deck and are not main features. The drawn cyan/navy edge colouring was recomputed from the committed masks with main's definition and reproduces si_border_pore exactly.
- Alignment: gr_chord_h_um = {N.v('ex_gr_chord_h_um'):.3f} um, gr_chord_v_um = {N.v('ex_gr_chord_v_um'):.3f} um, ratio gr_chord_ratio_hv = {N.v('ex_gr_chord_ratio_hv'):.3f} (features.csv; mean over ALL runs in the photo — the drawn lines are a few sample runs from the committed grey mask).

ASSUMPTIONS AND LIMITS
- Why this photo: {N.d['ex_id']['how']}. Other photos differ.
- Particle sizes and contact are 2-D section statistics; watershed splitting changes counts and sizes.
- Annotations (outline, circle, arrows, run lines) are editable PowerPoint shapes; tints are baked into the image.

SOURCES
- {repo_link(N, 'analysis/features.py', r'^def compute\(')}
- si_border_pore definition: {repo_link(N, 'analysis/features.py', r'^F\("si_border_pore"')}
- ring porosity definition: {repo_link(N, 'analysis/features.py', r'^F\("ring_porosity_250nm"')}

LIKELY QUESTION
Q: "Only 1.7% of the edge touches a pore, but 23% of particles touch one — which is right?"
A: Both. Most particles that touch a pore do so along a tiny part of their edge, so the share of edge is small while the share of particles is larger. That is exactly why we never report "contact" without saying which quantity it is.
{provenance_block(N)}""")
    return s


SLOTS, GAP_SLOTS = 17, 4


def strip_chart(slide, x, y, w, h, groups, ymin, ymax, ystep, yfmt, ytitle,
                name="chart"):
    """Editable strip chart: one dot per photo, one slot per photo inside a
    batch group, a median line per batch. Built as a line chart (markers,
    no connecting lines, gaps for blanks) because LibreOffice drops all but
    one series of python-pptx multi-series XY charts.
    groups: list of (batch, values, flagged_bool_list)."""
    import random
    from pptx.chart.data import CategoryChartData
    ncat = len(groups) * SLOTS + (len(groups) - 1) * GAP_SLOTS
    cd = CategoryChartData()
    cd.categories = [str(i + 1) for i in range(ncat)]
    styles = []
    for gi, (b, vals, flags) in enumerate(groups):
        off = gi * (SLOTS + GAP_SLOTS)
        n = len(vals)
        slots = [round(k * (SLOTS - 1) / max(n - 1, 1)) for k in range(n)] \
            if n > 1 else [SLOTS // 2]
        order = list(range(n))
        random.Random(7 + gi).shuffle(order)
        for want_flag in (False, True):
            col = [None] * ncat
            hit = False
            for k, i in enumerate(order):
                if bool(flags[i]) == want_flag:
                    col[off + slots[k]] = float(vals[i])
                    hit = True
            if hit:
                cd.add_series(f"{S.BATCH_SHORT[b]}{' (flagged)' if want_flag else ''}",
                              col)
                styles.append(("dot", b, want_flag))
        med = sorted(vals)[n // 2] if n % 2 else \
            (sorted(vals)[n // 2 - 1] + sorted(vals)[n // 2]) / 2
        col = [None] * ncat
        for k in range(3, SLOTS - 3):
            col[off + k] = float(med)
        cd.add_series(f"{S.BATCH_SHORT[b]} median", col)
        styles.append(("median", b, False))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x),
                                Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    ch = gf.chart
    ch.has_legend = False
    ch.font.name = S.BODY_FONT
    ch.font.size = Pt(12)
    ch.font.color.rgb = rgb(S.NAVY2)
    for (kind, b, fl), ser in zip(styles, ch.plots[0].series):
        ser.smooth = False
        if kind == "median":
            ser.marker.style = XL_MARKER_STYLE.NONE
            ser.format.line.color.rgb = rgb(S.NAVY)
            ser.format.line.width = Pt(2.5)
        else:
            ser.format.line.fill.background()
            ser.marker.style = XL_MARKER_STYLE.DIAMOND if fl else \
                XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 10 if fl else 8
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = rgb(S.BATCH_LIGHT[b] if fl
                                                        else S.BATCH[b])
            ser.marker.format.line.color.rgb = rgb(S.BATCH[b])
            ser.marker.format.line.width = Pt(1.25 if fl else 0.75)
    cs = ch._chartSpace.chart
    db = cs.find(qn("c:dispBlanksAs"))
    if db is None:
        db = etree.SubElement(cs, qn("c:dispBlanksAs"))
    db.set("val", "gap")
    va, xa = ch.value_axis, ch.category_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = ymin, ymax, ystep
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("E3E8EE")
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.color.rgb = rgb(S.NAVY2)
    va.tick_labels.font.size = Pt(12)
    va.tick_labels.number_format = yfmt
    va.tick_labels.number_format_is_linked = False
    va.major_tick_mark = XL_TICK_MARK.OUTSIDE
    va.has_title = True
    va.axis_title.text_frame.text = ytitle
    run_fmt(va.axis_title.text_frame.paragraphs[0].runs[0], size=12.5,
            color=S.NAVY)
    xa.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    xa.major_tick_mark = XL_TICK_MARK.NONE
    xa.format.line.color.rgb = rgb(S.NAVY2)
    xa.has_major_gridlines = False
    vax = ch._chartSpace.find(".//" + qn("c:valAx"))
    cb = vax.find(qn("c:crossBetween"))
    if cb is None:          # schema: crosses, crossBetween, majorUnit, ...
        cb = etree.Element(qn("c:crossBetween"))
        anchor = vax.find(qn("c:crosses"))
        if anchor is None:
            anchor = vax.find(qn("c:crossesAt"))
        anchor.addnext(cb)
    cb.set("val", "between")
    pa = cs.plotArea
    lay = pa.find(qn("c:layout"))
    if lay is None:
        lay = etree.Element(qn("c:layout"))
        pa.insert(0, lay)
    for c_ in list(lay):
        lay.remove(c_)
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    PL = dict(x=0.2, y=0.04, w=0.77, h=0.86)
    for tag, val in (("c:layoutTarget", "inner"), ("c:xMode", "edge"),
                     ("c:yMode", "edge"), ("c:x", PL["x"]), ("c:y", PL["y"]),
                     ("c:w", PL["w"]), ("c:h", PL["h"])):
        etree.SubElement(ml, qn(tag)).set("val", str(val))
    centres = []
    for gi in range(len(groups)):
        mid = gi * (SLOTS + GAP_SLOTS) + (SLOTS - 1) / 2 + 0.5
        centres.append(x + w * (PL["x"] + PL["w"] * mid / ncat))
    return ch, PL, centres


def slide6(prs, N, L):
    import pandas as pd
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "The key measurements, photo by photo")
    text(s, ML, 1.13, CW, 0.42, [dict(runs=[(
        "One dot per photo (one imaged location) · lines = batch medians · "
        "light diamonds = flagged by the quality guard",
        F(17, color=S.NAVY2))])])
    D = pd.read_csv(os.path.join(DECK, "provenance", "slide6_data.csv"))
    order = ["Batch_1", "Batch_2", "Batch_3"]
    cw_, gap = (CW - 2 * 0.3) / 3, 0.3
    y, h = 2.0, 2.95
    panels = [
        ("A · Visible pore space", "pore_frac", 100, 6, 18, 2, '0"%"',
         "pore share of photo area"),
        ("B · Bright-particle share", "silicon_frac", 100, 0, 20, 4, '0"%"',
         "bright share of photo area"),
        ("C · Image noise level", "noise_mad", 1, 6, 13, 1, "0",
         "noise (grey levels)"),
    ]
    for i, (pt, col, scale, y0, y1, st, fmt, yt) in enumerate(panels):
        x = ML + i * (cw_ + gap)
        text(s, x, y - 0.42, cw_, 0.36,
             [dict(runs=[(pt, F(15.5, bold=True))])], anchor="b")
        groups = [(b, (D[D.batch == b][col] * scale).tolist(),
                   D[D.batch == b].flagged.tolist()) for b in order]
        ch, PL, cen = strip_chart(s, x, y, cw_, h, groups, y0, y1, st, fmt,
                                  yt, name=pt)
        yb = y + h * (PL["y"] + PL["h"]) + 0.03
        for b, cx in zip(order, cen):
            text(s, cx - 0.62, yb, 1.24, 0.48, [
                dict(runs=[(S.BATCH_SHORT[b], F(12.5, bold=True,
                                                color=S.BATCH[b]))],
                     align="c"),
                dict(runs=[(f"n = {int((D.batch == b).sum())}",
                            F(11, color=S.GREY))], align="c")])
    ty = y + h + 0.55
    caps = [
        (f"Batch 1 sits lower on average ({d6(N, 'Batch_1', 'pore_frac')} "
         "vs Batch 3), but its photos overlap Batch 3's: a tentative shift."),
        (f"Two low-contrast Batch 1 photos read {N.t('si_flagged_values')}; "
         f"the other 29 lie between {N.t('si_range_unflagged')}. Without "
         "them, no batch differs."),
        ("Batch 1 and 2 photos are noisier than most Batch 3 photos. "
         "Texture, the one marker passing the corrected test, follows this "
         f"noise (r = {N.t('r_texture_noise')})."),
    ]
    for i, c in enumerate(caps):
        x = ML + i * (cw_ + gap)
        line(s, x, ty - 0.03, x + 0.6, ty - 0.03, color=S.TEAL, lw=1.5)
        text(s, x, ty + 0.02, cw_, 0.95, [dict(runs=[(c, F(13.5))],
                                              ls=0.95)])
    rect(s, ML, 6.48, CW, 0.52, fill=S.PANEL)
    text(s, ML + 0.15, 6.48, CW - 0.3, 0.52, [dict(runs=[
        ("Imaging is entangled with batch: ", F(13.5, bold=True,
                                                color=S.AMBER)),
        (f"{N.t('n_B3_lower_noise')} of the {N.t('n_Batch_3')} Batch 3 photos "
         "were taken at lower noise than any Batch 1 or 2 photo, so batch "
         "and microscope-settings differences cannot be fully separated.",
         F(13.5))])], anchor="m")
    L.add(6, "Batch 1 pore space -14% vs Batch 3 (suggestive)",
          "d_Batch_1_pore_frac_pct", "d_Batch_1_pore_frac_class")
    L.add(6, "Flagged Batch 1 photos read 16.9% and 18.4% bright; others "
          "4.5-9.7%", "si_flagged_values", "si_range_unflagged")
    L.add(6, "Texture is the only Holm-significant marker (B1 +18%, B2 +15%)",
          "holm_survivors", "d_Batch_1_bse_bulk_texture_pct",
          "d_Batch_2_bse_bulk_texture_pct")
    L.add(6, "Texture correlates with image noise, r = 0.91 (31 photos)",
          "r_texture_noise")
    L.add(6, "12 of 17 Batch 3 photos have lower noise than any Batch 1/2 "
          "photo", "n_B3_lower_noise")
    footer(s, 6, N)
    nl = N.v("noise_levels")
    notes(s, f"""
WHAT TO SAY
Three panels, one dot per photo. A: visible pore space — Batch 1 is a little lower on average, but its dots overlap Batch 3's, which is why we call it tentative. B: bright-particle share — the two light diamonds are the photos the quality guard flagged; they alone create Batch 1's apparent +45%; the other 29 photos all sit in the same range. C: image noise — every Batch 1 and Batch 2 photo is at the upper noise levels, while most Batch 3 photos are cleaner. The one marker that survives the strict test, grey-level texture, rises with exactly this noise. That is the most important caveat on this dataset: some batch differences may be microscope-settings differences.

TERMS
- Median line: the middle value of the batch's photos.
- Image noise: main's quality guard noise_mad — the spread of fine pixel-to-pixel fluctuations inside the grey bulk (median absolute deviation x 1.4826, in grey levels). It takes only a few distinct values: {', '.join(f"{b.replace('_', ' ')} {v}" for b, v in nl.items())}.
- Grey-bulk texture: bse_bulk_texture — standard deviation of BSE grey level inside the grey (graphite) class.
- r: Pearson correlation over the 31 photos.

HOW THE NUMBERS WERE OBTAINED
- Data: features.csv (subset == full) joined to quality_guards.csv; exported as provenance/slide6_data.csv. The charts are native, editable PowerPoint charts built from that file (right-click > Edit Data). Each photo has its own slot inside its batch group (order shuffled with a fixed seed), so the horizontal position inside a batch has no meaning; blank cells are gaps.
- Panel A: Batch 1 mean shift {N.pct('d_Batch_1_pore_frac_pct', 1, signed=True)} ({N.t('d_Batch_1_pore_frac_class')}; interval {N.t('d_Batch_1_pore_frac_ci')}, Holm p {N.t('d_Batch_1_pore_frac_pholm')}).
- Panel B: flagged photos {N.t('flagged_ids')} read {N.t('si_flagged_values')}; unflagged range {N.t('si_range_unflagged')}. Batch 1 shift {d6(N, 'Batch_1', 'silicon_frac')} with flagged photos, {d6(N, 'Batch_1', 'silicon_frac', True)} without (both undetermined).
- Panel C: noise_mad per photo (quality_guards.csv). The texture link: r(bse_bulk_texture, noise_mad) = {N.t('r_texture_noise')}; texture shift Batch 1 {N.pct('d_Batch_1_bse_bulk_texture_pct', signed=True)} (Holm p {N.t('d_Batch_1_bse_bulk_texture_pholm')}), Batch 2 {N.pct('d_Batch_2_bse_bulk_texture_pct', signed=True)} (Holm p {N.t('d_Batch_2_bse_bulk_texture_pholm')}). Main itself labels bse_bulk_texture and etd_roughness "artefact-risk" (|r| with noise > 0.7) and excludes them from its artefact-safe classifier. Other noise correlations: inlens_edge_density r = {N.t('r_noise_inlens_edge_density')}, etd_roughness r = {N.t('r_noise_etd_roughness')}, pore_frac r = {N.t('r_noise_pore_frac')}.
- Observation unit: one photo = one imaged location; n = {N.t('n_Batch_1')} / {N.t('n_Batch_2')} / {N.t('n_Batch_3')}. No error bars are drawn because the photos within a batch may not be independent samples of the batch.

ASSUMPTIONS AND LIMITS
- We do not know how the locations were chosen or whether some photos come from the same piece of electrode; frame heights also differ ({N.t('frame_heights_px')}), which hints at different acquisition sessions (not analysed on main).
- The noise measure is coarse (4 distinct values), so panel C shows dots on a few levels only.

SOURCES
- {repo_link(N, 'analysis/segment.py', r'^def quality_guards\(')}
- Artefact-risk rule: {repo_link(N, 'analysis/categorise.py', r'^ARTEFACT_RISK')}

LIKELY QUESTION
Q: "So is Batch 1 different or not?"
A: Not established. Its pore space is slightly lower on average, but that does not survive the correction for testing 20 markers, the photos overlap, and Batch 1 was photographed under noisier conditions. The honest call is "tentative — re-image under matched settings".
{provenance_block(N)}""")
    return s


def d6(N, b, f, nf=False):
    return N.pct(f"d_{'nf_' if nf else ''}{b}_{f}_pct", signed=True)


def est_cell_h(paras, width):
    h = 0.1
    for p in paras:
        if isinstance(p, tuple):
            p = dict(runs=[p])
        txt = "".join(t_ for t_, _ in p["runs"])
        if not txt.strip():
            continue
        pt = max(f.get("size", 20) for _, f in p["runs"])
        bold = all(f.get("bold") for _, f in p["runs"])
        ind = p.get("indent", 0.22) if p.get("bullet") else 0
        nl = n_lines(txt, pt, width - 0.14 - ind, bold)
        h += nl * pt * 1.17 * p.get("ls", 1.0) / 72 + p.get("sb", 0) / 72
    return h


def slide7(prs, N, L, newb):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Results by batch — and the next decision")

    def chip(t_, col, sub=None):
        out = [dict(runs=[(t_, F(13.5, bold=True, color=col))], ls=0.9)]
        if sub:
            out.append(dict(runs=[(sub, F(11.5, color=S.GREY))], ls=0.9,
                            sb=2))
        return out

    def obs(*items):
        return [dict(runs=[(it, F(12.5))], bullet="•", ls=0.9,
                     sb=0 if i == 0 else 2) for i, it in enumerate(items)]

    def follow(t_):
        return [dict(runs=[(t_, F(12.5, color=S.NAVY2))], ls=0.9)]

    def batch_cell(b, n, sub):
        out = [dict(runs=[(S.BATCH_SHORT.get(b, b),
                           F(14, bold=True, color=S.BATCH.get(b, S.NAVY)))])]
        if n:
            out.append(dict(runs=[(f"{n} photos", F(11.5, color=S.GREY))]))
        if sub:
            out.append(dict(runs=[(sub, F(11, italic=True, color=S.GREY))]))
        return out

    rows = [
        (batch_cell("Batch_1", N.t("n_Batch_1"), ""),
         obs(f"Less visible pore space ({d6(N, 'Batch_1', 'pore_frac')}) and "
             f"fewer large gaps ({d6(N, 'Batch_1', 'large_gap_frac')})",
             f"Bright share {d6(N, 'Batch_1', 'silicon_frac')} due to two "
             f"low-contrast photos ({d6(N, 'Batch_1', 'silicon_frac', True)} "
             "without them)",
             f"Texture {d6(N, 'Batch_1', 'bse_bulk_texture')}, which tracks "
             "image noise"),
         chip("Tentative", S.AMBER, "no material shift survives the "
              "correction"),
         follow("Re-image the two flagged locations; add 5+ locations with "
                "Batch 3's microscope settings")),
        (batch_cell("Batch_2", N.t("n_Batch_2"), ""),
         obs("No established structural shift: finest pores "
             f"{d6(N, 'Batch_2', 'pore_thick_d10_um')} and flake texture order "
             f"{d6(N, 'Batch_2', 'gr_st_coherence')} are tentative",
             f"Texture {d6(N, 'Batch_2', 'bse_bulk_texture')}, which tracks "
             "image noise"),
         chip("No difference established", S.TEAL),
         follow("No action on structure; keep imaging settings matched to "
                "Batch 3")),
        (batch_cell("Batch_3", N.t("n_Batch_3"), "reference"),
         obs("Defines the normal range for every marker",
             f"Held out, {N.t('b3_out')} of its own photos fall outside its "
             "single-photo envelope"),
         chip("Reference", S.NAVY2, "envelope too tight for single photos"),
         follow("Judge groups of photos; widen the reference with more "
                "known-good batches")),
    ]
    if newb:
        obs_t = newb["observations"]
        if len(obs_t) > 330:
            obs_t = obs_t[:327].rsplit(" ", 1)[0] + " …"
        st_head, _, st_sub = newb["status"].partition(" — ")
        rows.append((
            batch_cell(newb["name"], newb.get("n_images"), "new delivery"),
            [dict(runs=[(obs_t, F(12.5))], ls=0.9)],
            chip(st_head, S.AMBER, st_sub or None),
            follow(newb["follow_up"])))
    else:
        rows.append((
            batch_cell("New batch", None, "results being generated"),
            [dict(runs=[("Pending — the run on the new delivery is in "
                         "progress; its results are inserted here.",
                         F(12.5, italic=True, color=S.GREY))], ls=0.9)],
            chip("Pending", S.GREY),
            follow("Write down the expected answer before the results are "
                   "revealed")))
    colw = [1.62, 5.28, 2.3, 3.033]
    rh = [0.4] + [max(0.5, max(est_cell_h(c, w_) for c, w_ in zip(r, colw)))
                  for r in rows]
    tbl = booktabs(s, ML, 1.22, colw, rh,
                   ["Batch", "Main measured observations", "Status",
                    "Follow-up justified by the evidence"], rows,
                   name="results table")
    if not newb:
        for j in range(4):
            c = tbl.cell(len(rows), j)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb("F6F7F9")
    L.add(7, "Batch 1: pore space and large-gap shifts (suggestive)",
          "d_Batch_1_pore_frac_pct", "d_Batch_1_large_gap_frac_pct",
          "d_Batch_1_large_gap_frac_class")
    L.add(7, "Batch 1 bright share +45% with flagged photos, -2% without",
          "d_Batch_1_silicon_frac_pct", "d_nf_Batch_1_silicon_frac_pct")
    L.add(7, "Texture +18% (B1) / +15% (B2), Holm-significant, noise-linked",
          "d_Batch_1_bse_bulk_texture_pct", "d_Batch_2_bse_bulk_texture_pct",
          "d_Batch_1_bse_bulk_texture_pholm", "r_texture_noise")
    L.add(7, "Batch 2 graphite texture order -8% and finest pores -7% "
          "(both suggestive)", "d_Batch_2_gr_st_coherence_pct",
          "d_Batch_2_gr_st_coherence_class", "d_Batch_2_pore_thick_d10_um_pct",
          "d_Batch_2_pore_thick_d10_um_class")
    L.add(7, "7 of 17 held-out Batch 3 photos called OUT of the Batch 3 "
          "envelope", "b3_out")
    if CR:
        L.add(7, "From-scratch rebuild: same structural results; classifier "
              f"LOO all-features {CR['loo_allfeatures_fresh'][0]}/31, "
              f"artefact-safe {CR['loo_safe_fresh'][0]}/31, B3 OUT "
              f"{CR['b3_out_fresh'][0]}/17 (notes only)",
              note="provenance/clean_rebuild_summary.json — main's run_all.py "
                   "--fast on an emptied outputs/ folder, Python 3.10 env")
    if newb:
        L.add(7, f"{newb['name']}: {newb['observations']}",
              note=f"inserted from {newb.get('source')} "
                   f"({newb.get('source_kind')}); not verified by this deck")
    yt = 1.22 + sum(rh) + 0.14
    th = 0.74
    room = 7.02 - yt          # space left above the footer
    show_sim = room >= th + 0.1 + 0.6 + 0.48
    tight = room < th + 0.1 + 0.6
    rect(s, ML, yt, CW, th, fill=S.TEAL_LIGHT, name="take-home")
    rect(s, ML, yt, 0.06, th, fill=S.TEAL)
    text(s, ML + 0.22, yt, CW - 0.35, th, [dict(runs=[
        ("Take-home: ", F(16, bold=True, color=S.TEAL)),
        ("the pipeline measures structure reproducibly and catches unreliable "
         "photos. The batches look broadly alike — Batch 1's lower pore space "
         "is tentative — and single-photo batch calls are not yet reliable.",
         F(16))], ls=0.95)], anchor="m")
    yn = yt + th + 0.1
    nv = 12.5 if tight else 13.5
    text(s, ML, yn, CW, 0.6, [dict(runs=[
        ("Most useful next validation: ", F(nv, bold=True)),
        ("a blind test — photograph 5 or more locations of the new delivery with "
         "Batch 3's microscope settings, write down the expected answer "
         "first, have an expert fill main's 300-point check sheet, and "
         "confirm the bright particles chemically (EDS).",
         F(nv, color=S.NAVY2))], ls=0.95)])
    sim_box = text if show_sim else (lambda *a_, **k_: None)
    sim_box(s, ML, yn + 0.62, CW, 0.44, [dict(runs=[
        ("Simulation branch (model scenarios, not measurements): ",
         F(11, bold=True, color=S.GREY)),
        (f"at 2C with measured porosity, Batch 1/2 run "
         f"{N.pct('dfn_Batch_1_measured_2C', 1, signed=True)}/"
         f"{N.pct('dfn_Batch_2_measured_2C', 1, signed=True)} vs Batch 3; "
         "with porosity set to 30% they match; at 1C and slower every run stops at "
         f"the simulation time limit ({N.t('dfn_censored_ah')}).",
         F(11, color=S.GREY))], ls=0.92)])
    L.add(7, "DFN 2C, measured porosity: B1 -2.7%, B2 -2.9% vs B3; "
          "corrected porosity: no difference", "dfn_Batch_1_measured_2C",
          "dfn_Batch_2_measured_2C", "dfn_Batch_1_corrected_2C",
          "dfn_Batch_2_corrected_2C")
    L.add(7, "DFN at <=1C: every run returns 5.25 Ah = time cut-off",
          "dfn_censored_ah", "dfn_censored_spread")
    footer(s, 7, N)
    nb_notes = ""
    if newb:
        nb_notes = (f"\n\nNEW BATCH ROW (inserted from {newb.get('source')}; "
                    f"source type: {newb.get('source_kind')})\n"
                    f"{newb.get('notes', '')}\n"
                    "Treat this call as indicative: main's group-of-photos "
                    "accuracy is measured in-sample (see below), and the "
                    "single-photo envelope rejects "
                    f"{N.t('b3_out')} known-good Batch 3 photos.")
    notes(s, f"""
WHAT TO SAY
Batch by batch. Batch 1 shows the most movement — less visible pore space and fewer large gaps — but nothing material survives the correction for testing 20 markers, its bright-particle "excess" comes from two poor photos, and its texture shift follows image noise. So: tentative, re-image under matched settings. Batch 2 shows no established structural shift; its only significant change is again the noise-linked texture. Batch 3 is the reference, and one honest finding is that its own envelope is too tight for single photos: hold one Batch 3 photo out and it is called "out of spec" {N.t('b3_out')} times. The last row is the new delivery we are running now. Take-home: the measurement part works and is reproducible; the decision part needs a blind test on new, properly matched photos.

TERMS
- Envelope: the range of "normal" for Batch 3 — a robust distance from the Batch 3 average with a cut-off at the 95th percentile of Batch 3's own photos.
- Held out (leave-one-out): the photo being judged is removed before the reference is built.
- Blind test: the expected answer is written down before the result is seen, so nobody tunes the method to the answer.
- EDS: energy-dispersive X-ray spectroscopy, a chemical map that could confirm the bright particles are silicon.

HOW THE NUMBERS WERE OBTAINED
- Batch rows: deltas.csv / deltas_noflag.csv (kept markers; Batch 1 and 2 each vs Batch 3). Large gaps = large_gap_frac (share of pore area in pores each bigger than 0.5% of the photo): Batch 1 {d6(N, 'Batch_1', 'large_gap_frac')} ({N.t('d_Batch_1_large_gap_frac_class')}; highly variable between photo halves). Holm-significant rows: {N.t('holm_survivors')} only.
- Graphite texture order = gr_st_coherence (structure-tensor coherence inside the grey class): Batch 2 {d6(N, 'Batch_2', 'gr_st_coherence')} ({N.t('d_Batch_2_gr_st_coherence_class')}). Finest pores = pore_thick_d10_um (10th percentile of pore local thickness): Batch 2 {d6(N, 'Batch_2', 'pore_thick_d10_um')} ({N.t('d_Batch_2_pore_thick_d10_um_class')}; this marker sits on the 2-pixel resolution floor, 0.05 um).
- Envelope check: loo_predictions_allfeatures.csv, call_inout for the 17 Batch 3 photos ({N.t('b3_out')} OUT). Note: the committed loo_summary.csv "false_alarm" column divides by all 31 photos and is stale; regenerated summary: provenance/loo_summary_regenerated.csv.
- Batch calls: single photos {N.t('loo_correct')}/31 correct ({N.t('loo_acc')}; Batch 1 {N.t('loo_Batch_1')}, Batch 2 {N.t('loo_Batch_2')}, Batch 3 {N.t('loo_Batch_3')}); always-"Batch 3" {N.t('majority_acc')}; shuffled labels {N.t('shuffled_acc')}; artefact-safe variant {N.t('loo_acc_safe')}. Main's README/report quote "~45% vs 37% shuffled" — from the stale summary. Main's "groups of 5+ photos reach ~97-100%" (images_needed.csv: Batch 1 {N.t('group5_Batch_1')}, Batch 2 {N.t('group5_Batch_2')} at n = 5) is IN-SAMPLE: the centroids are fitted on the same photos that are sampled, so it is not an independent validation.
- Simulation (dfn_capacity.csv, dfn_deltas.csv): at 2C, measured porosity, mean capacity {N.t('dfn_cap2C_Batch_1')} / {N.t('dfn_cap2C_Batch_2')} / {N.t('dfn_cap2C_Batch_3')} (B1/B2/B3) -> {N.pct('dfn_Batch_1_measured_2C', 1, signed=True)} and {N.pct('dfn_Batch_2_measured_2C', 1, signed=True)}; the quoted intervals are re-sampling ranges over 30 runs, i.e. scenario spread under fixed assumptions, not confidence intervals for a real cell. With porosity set to 0.30 the batches agree within 0.01%. At C/10, C/2 and 1C every run returns {N.t('dfn_censored_ah')} with spread {N.t('dfn_censored_spread')} Ah — the simulation stops at its time limit (dfn.py: t_end = 1.05 x 3600 s / C-rate), so those rates carry no comparison.

ASSUMPTIONS AND LIMITS
- ACCEPT / INVESTIGATE / REJECT labels are not used: main has no validated decision rule of that kind. Status words here describe the evidence, not product quality.
- No quality ranking is implied by batch names.
- With 7 photos per batch, "no difference established" does not mean "identical".

SOURCES
- {repo_link(N, 'analysis/deltas.py')}
- {repo_link(N, 'analysis/categorise.py', r'^def images_needed\(')} (in-sample group calibration)
- {repo_link(N, 'analysis/dfn.py', r't_end = ')}
- Webapp (same code, artefact-safe 6-feature model): {repo_link(N, 'webapp/app.py')}

LIKELY QUESTION
Q: "Should we reject Batch 1?"
A: The evidence does not support a rejection. It supports re-imaging: two of its seven photos are unreliable, its remaining shifts are tentative, and it was photographed under noisier conditions than Batch 3. A matched, blind re-test is cheap and would settle it.{nb_notes}{rebuild_caveat(N)}
{provenance_block(N)}""")
    return s


S.TEAL_RUN = "0E7C86"
S.VRUN = "C25E00"


# =============================================================================
def presenter_script(prs, path, N, name=OUT_NAME):
    out = ["# Presenter script — TeamJester algorithm overview", "",
           f"Deck: `{name}.pptx` · source: GitHub main @ "
           f"`{N.v('sha')}` ({N.t('commit_date')}). "
           "Read the WHAT TO SAY part of each slide; the full speaker notes "
           "(definitions, number provenance, limits, Q&A) are in the deck.", ""]
    caveat = ""
    for i, sl in enumerate(prs.slides, start=1):
        nt = sl.notes_slide.notes_text_frame.text
        say = nt.split("WHAT TO SAY", 1)[1].split("\n\nTERMS", 1)[0].strip()
        q = nt.split("LIKELY QUESTION", 1)[1].split("\n\nPROVENANCE", 1)[0]
        if "BUILD-STATE CAVEAT" in q:
            q, cv = q.split("BUILD-STATE CAVEAT", 1)
            caveat = "BUILD-STATE CAVEAT" + cv
        nb = ""
        if "NEW BATCH ROW" in q:
            q, nb = q.split("NEW BATCH ROW", 1)
            nb = "NEW BATCH ROW" + nb
        q = q.strip()
        ttl = [sh for sh in sl.shapes if sh.name == "Title"][0].text_frame.text
        out += [f"## {i}. {ttl}", "", say, ""]
        if nb:
            out += ["**New batch row:** " + nb.strip().split("\n", 1)[0], "",
                    "```", nb.strip().split("\n", 1)[1].strip(), "```", ""]
        out += ["**If asked:**", "", q.replace("\nA:", "\n\nA:"), ""]
    if caveat:
        head, body = caveat.strip().split("\n", 1)
        out += ["## Background for questions on reproducibility", "",
                f"*{head}*", "", body.strip(), ""]
    with open(path, "w") as f:
        f.write("\n".join(out))


def export_pdf(pptx, outdir):
    """LibreOffice headless, with an isolated profile whose user/fonts holds
    the Latin Modern fonts (headless soffice does not see ~/Library/Fonts)."""
    prof = os.path.join(DECK, ".lo_profile")
    fdir = os.path.join(prof, "user", "fonts")
    os.makedirs(fdir, exist_ok=True)
    src = os.path.expanduser("~/Library/Fonts")
    for fn in os.listdir(src) if os.path.isdir(src) else []:
        if re.match(r"lm(sans10|roman1[027]|mono10)-.*\.otf$", fn) and \
                not os.path.exists(os.path.join(fdir, fn)):
            shutil.copy(os.path.join(src, fn), fdir)
    cmd = ["soffice", f"-env:UserInstallation=file://{prof}", "--headless",
           "--convert-to", "pdf", "--outdir", outdir, pptx]
    subprocess.run(cmd, check=True, capture_output=True, timeout=300)
    return os.path.join(outdir, os.path.basename(pptx)[:-5] + ".pdf")


def render_png(pdf, outdir, dpi=110):
    import fitz
    os.makedirs(outdir, exist_ok=True)
    doc = fitz.open(pdf)
    paths = []
    for i, page in enumerate(doc, start=1):
        p = os.path.join(outdir, f"slide-{i}.png")
        page.get_pixmap(dpi=dpi).save(p)
        paths.append(p)
    fonts = sorted({f[3] for page in doc for f in page.get_fonts()})
    return paths, fonts, len(doc)


def main():
    global REPO
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True)
    ap.add_argument("--new-batch", default=None,
                    help="categorise_results.csv, webapp JSON or manual JSON")
    ap.add_argument("--new-batch-name", default=None)
    ap.add_argument("--out", default=OUT_NAME,
                    help="output basename (default %(default)s); written "
                         "next to this README")
    ap.add_argument("--outdir", default=DECK)
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()
    REPO = os.path.abspath(a.repo)
    N = Nums()
    vp = os.path.join(DECK, "provenance", "verify_summary.json")
    verify = json.load(open(vp)) if os.path.exists(vp) else {}
    ve = os.path.join(DECK, "provenance", "verify_environment.json")
    if os.path.exists(ve) and json.load(open(ve)).get("head") != N.v("sha"):
        print("WARNING: the worktree is not at the commit that "
              "src/verify_main.py verified — re-run ./build.sh --verify "
              "before presenting.", file=sys.stderr)
    cr = os.path.join(DECK, "provenance", "clean_rebuild_summary.json")
    if os.path.exists(cr):
        CR.update(json.load(open(cr)))
        verify["clean_rebuild"] = CR.get("tables", {})
    with open(os.path.join(A, "annotations.json")) as f:
        ann = json.load(f)
    newb = NB.load(a.new_batch, a.new_batch_name) if a.new_batch else None
    L = Ledger(N, verify, N.v("sha"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = "TeamJester — explainable SEM batch comparison"
    prs.core_properties.subject = (f"GitHub main @ {N.v('sha')} "
                                   f"({N.t('commit_date')})")
    prs.core_properties.author = "TeamJester"
    slide1(prs, N, L, ann)
    slide2(prs, N, L)
    slide3(prs, N, L, ann)
    slide4(prs, N, L)
    slide5(prs, N, L, ann)
    slide6(prs, N, L)
    slide7(prs, N, L, newb)
    os.makedirs(a.outdir, exist_ok=True)
    pptx = os.path.join(a.outdir, a.out + ".pptx")
    prs.save(pptx)
    sfx = "" if a.out == OUT_NAME else f"_{a.out}"
    L.write(os.path.join(a.outdir, f"CLAIM_LEDGER{sfx}.csv"))
    presenter_script(prs, os.path.join(a.outdir, f"PRESENTER_SCRIPT{sfx}.md"),
                     N, a.out)
    print("wrote", pptx, f"({len(L.rows)} ledger rows)")
    if not a.no_pdf:
        pdf = export_pdf(pptx, a.outdir)
        pngs, fonts, n = render_png(pdf, os.path.join(
            a.outdir, "previews" + sfx))
        bad = [f for f in fonts if not re.search(r"LM(Sans|Roman|Mono)", f)]
        print("wrote", pdf, f"{n} pages; non-Latin-Modern fonts: {bad}")


if __name__ == "__main__":
    main()
