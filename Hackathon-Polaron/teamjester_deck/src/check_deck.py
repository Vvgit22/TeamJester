"""QA for the built deck: schema element order, provenance on every slide,
PPTX <-> PDF text agreement, fonts, and hyperlinks.

    python src/check_deck.py [--deck TeamJester_Algorithm_Overview] [--offline]

Exit code 1 if any check fails.
"""
import argparse
import json
import os
import re
import subprocess
import sys
import zipfile

from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.dirname(HERE)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
ORDER = {
    f"{{{A}}}pPr": ["lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx",
                    "buSzPct", "buSzPts", "buFontTx", "buFont", "buNone",
                    "buAutoNum", "buChar", "buBlip", "tabLst", "defRPr",
                    "extLst"],
    f"{{{A}}}rPr": ["ln", "noFill", "solidFill", "gradFill", "blipFill",
                    "pattFill", "grpFill", "effectLst", "effectDag",
                    "highlight", "uLnTx", "uLn", "uFillTx", "uFill", "latin",
                    "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl",
                    "extLst"],
    f"{{{A}}}ln": ["noFill", "solidFill", "gradFill", "pattFill", "prstDash",
                   "custDash", "round", "bevel", "miter", "headEnd", "tailEnd",
                   "extLst"],
    f"{{{A}}}tcPr": ["lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr",
                     "cell3D", "noFill", "solidFill", "gradFill", "blipFill",
                     "pattFill", "grpFill", "headers", "extLst"],
    f"{{{P}}}spPr": ["xfrm", "custGeom", "prstGeom", "noFill", "solidFill",
                     "gradFill", "blipFill", "pattFill", "grpFill", "ln",
                     "effectLst", "effectDag", "scene3d", "sp3d", "extLst"],
    f"{{{C}}}valAx": ["axId", "scaling", "delete", "axPos", "majorGridlines",
                      "minorGridlines", "title", "numFmt", "majorTickMark",
                      "minorTickMark", "tickLblPos", "spPr", "txPr",
                      "crossAx", "crosses", "crossesAt", "crossBetween",
                      "majorUnit", "minorUnit", "dispUnits", "extLst"],
    f"{{{C}}}chart": ["title", "autoTitleDeleted", "pivotFmts", "view3D",
                      "floor", "sideWall", "backWall", "plotArea", "legend",
                      "plotVisOnly", "dispBlanksAs", "showDLblsOverMax",
                      "extLst"],
}


def check_order(zf):
    errs = []
    for name in zf.namelist():
        if not name.endswith(".xml") or not (name.startswith("ppt/slides/")
                                            or name.startswith("ppt/charts/")):
            continue
        root = etree.fromstring(zf.read(name))
        for tag, order in ORDER.items():
            for el in root.iter(tag):
                kids = [etree.QName(k).localname for k in el
                        if isinstance(k.tag, str)]
                idx = [order.index(k) for k in kids if k in order]
                if idx != sorted(idx):
                    errs.append(f"{name}: <{etree.QName(tag).localname}> "
                                f"children out of order: {kids}")
                if tag.endswith("plotArea"):
                    pass
        for pa in root.iter(f"{{{C}}}plotArea"):
            first = etree.QName(pa[0]).localname
            if first != "layout":
                errs.append(f"{name}: plotArea does not start with layout")
    return errs


def slide_texts(pptx):
    from pptx import Presentation
    prs = Presentation(pptx)
    out = []
    for s in prs.slides:
        txt = []
        for sh in s.shapes:
            if sh.has_text_frame:
                txt.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        txt.append(c.text_frame.text)
        links = [r.hyperlink.address for sh in s.shapes if sh.has_text_frame
                 for p in sh.text_frame.paragraphs for r in p.runs
                 if r.hyperlink.address]
        for sh in s.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        links += [r.hyperlink.address for p in
                                  c.text_frame.paragraphs for r in p.runs
                                  if r.hyperlink.address]
        out.append(dict(text="\n".join(txt), links=links,
                        notes=s.notes_slide.notes_text_frame.text))
    return out


def norm_words(t):
    t = t.replace("\u2212", "-").replace("\u00a0", " ")
    return [w for w in re.findall(r"[A-Za-z0-9%.]+", t) if len(w) > 1]


def check_links(links, sha, offline):
    errs = []
    for url in sorted(set(links)):
        if offline:
            continue
        m = re.match(r"https://github.com/([^/]+/[^/]+)/blob/([0-9a-f]{40})/"
                     r"([^#]+)(?:#L(\d+))?$", url)
        if m:
            repo, ref, path, line = m.groups()
            if ref != sha:
                errs.append(f"{url}: not pinned to {sha[:7]}")
                continue
            r = subprocess.run(["gh", "api", f"repos/{repo}/contents/{path}"
                                f"?ref={ref}", "--jq", ".size"],
                               capture_output=True, text=True)
            if r.returncode != 0:
                errs.append(f"{url}: gh api failed ({r.stderr.strip()[:80]})")
            continue
        r = subprocess.run(["curl", "-sIL", "-o", "/dev/null", "-w",
                            "%{http_code}", "--max-time", "25", url],
                           capture_output=True, text=True)
        code = r.stdout.strip()[-3:]
        if code not in ("200", "301", "302", "303", "403", "418"):
            errs.append(f"{url}: HTTP {code}")
    return errs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default="TeamJester_Algorithm_Overview")
    ap.add_argument("--dir", default=DECK)
    ap.add_argument("--offline", action="store_true")
    a = ap.parse_args()
    pptx = os.path.join(a.dir, a.deck + ".pptx")
    pdf = os.path.join(a.dir, a.deck + ".pdf")
    nums = json.load(open(os.path.join(DECK, "provenance",
                                       "deck_numbers.json")))
    sha = nums["sha"]["value"]
    fails = []
    with zipfile.ZipFile(pptx) as zf:
        fails += check_order(zf)
    S = slide_texts(pptx)
    if len(S) != 7:
        fails.append(f"expected 7 slides, found {len(S)}")
    for i, s in enumerate(S, 1):
        for need in (sha, "github.com/Augustin-Briens/TeamJester",
                     "WHAT TO SAY", "TERMS", "HOW THE NUMBERS WERE OBTAINED",
                     "ASSUMPTIONS AND LIMITS", "SOURCES", "LIKELY QUESTION"):
            if need not in s["notes"]:
                fails.append(f"slide {i}: notes missing '{need[:40]}'")
        if f"main @ {sha[:7]}" not in s["text"] or f"{i} / 7" not in s["text"]:
            fails.append(f"slide {i}: footer (version/date or number) missing")
        for bad in ("nan", "None", "{", "}", "TODO"):
            if re.search(rf"(?<![A-Za-z]){re.escape(bad)}(?![A-Za-z])",
                         s["text"]):
                fails.append(f"slide {i}: suspicious token '{bad}' in text")
    import fitz
    doc = fitz.open(pdf)
    if len(doc) != len(S):
        fails.append(f"PDF has {len(doc)} pages, PPTX {len(S)} slides")
    fonts = sorted({f[3] for pg in doc for f in pg.get_fonts()})
    other = [f for f in fonts if not re.search(r"LM(Sans|Roman|Mono)", f)]
    if other:
        fails.append(f"PDF uses fallback fonts: {other}")
    for i, (s, pg) in enumerate(zip(S, doc), 1):
        pw = set(norm_words(pg.get_text()))
        miss = [w for w in norm_words(s["text"]) if w not in pw]
        if miss:
            fails.append(f"slide {i}: {len(miss)} PPTX words not found in "
                         f"PDF page: {miss[:8]}")
    links = [u for s in S for u in s["links"]]
    fails += check_links(links, sha, a.offline)
    print(f"slides: {len(S)} · PDF pages: {len(doc)} · hyperlinks: "
          f"{len(set(links))} · fonts: {len(fonts)} (all Latin Modern: "
          f"{not other})")
    if fails:
        print("FAIL:\n  " + "\n  ".join(fails))
        sys.exit(1)
    print("all checks passed")


if __name__ == "__main__":
    main()
